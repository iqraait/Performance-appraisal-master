import openpyxl
import xlrd
import csv
import io
import datetime
import re
from html.parser import HTMLParser
import xml.etree.ElementTree as ET
from django.shortcuts import render
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Count, Avg, Q
from django.contrib.auth.models import User
from rest_framework import viewsets, status, permissions
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.decorators import action
from rest_framework.authtoken.models import Token
from rest_framework.authtoken.views import ObtainAuthToken

from .models import (
    Employee, Appraisal, Branch, DepartmentAdmin, TrainingEmployee, 
    WhatsAppConfig, EmailConfig, TraineeAssessment,
    IncentivePeriod, DivisionRate, EmployeeLevelConfig, ExperienceLevelConfig, IncentiveRecord,
    DutyLocation
)
from .serializers import (
    EmployeeSerializer, AppraisalSerializer, UserSerializer, 
    BranchSerializer, DepartmentAdminSerializer, TrainingEmployeeSerializer, WhatsAppConfigSerializer,
    EmailConfigSerializer, TraineeAssessmentSerializer,
    IncentivePeriodSerializer, DivisionRateSerializer, EmployeeLevelConfigSerializer, 
    ExperienceLevelConfigSerializer, IncentiveRecordSerializer,
    DutyLocationSerializer
)
from .whatsapp_service import (
    send_trainee_assessment_reminder, check_and_send_all_due_reminders,
    send_meta_whatsapp_message, get_whatsapp_config, clean_phone_number
)
from .email_service import (
    send_trainee_assessment_email, check_and_send_all_due_email_reminders,
    send_department_wise_trainee_emails,
    get_department_emails, get_email_config, send_raw_email
)

# Custom permissions
class IsAdminUserOrReadOnly(permissions.BasePermission):
    def has_permission(self, request, view):
        if request.user and request.user.is_authenticated:
            return request.user.is_staff or request.user.is_superuser or hasattr(request.user, 'department_admin')
        return False


# Location matching helpers (strict exact full-name token matching)
def clean_loc_token(s):
    if not s:
        return ''
    return re.sub(r'^(?:na[\\/]+)', '', str(s).strip(), flags=re.IGNORECASE).strip()

def build_exact_location_filter(loc_str, field_name='location'):
    """
    Builds a Q filter that strictly matches the full location name as an exact token,
    whether standalone, prefixed with NA\\ or NA/, or as an item in a comma-separated list.
    E.g. 'WARD' matches 'WARD', 'NA\\WARD', 'WARD, OPD', but NEVER matches 'GENERAL WARD MALE'.
    """
    if not loc_str or not str(loc_str).strip():
        return Q()
    raw = str(loc_str).strip()
    clean = clean_loc_token(raw)
    tokens = {raw, clean} if clean else {raw}
    
    q = Q()
    for t in tokens:
        if not t:
            continue
        kw = lambda suffix: f'{field_name}__{suffix}'
        # 1. Exact standalone match (with or without NA\\ / NA/ prefix)
        q |= Q(**{kw('iexact'): t})
        q |= Q(**{kw('iexact'): f'NA\\{t}'})
        q |= Q(**{kw('iexact'): f'NA/{t}'})
        # 2. Comma-separated: at the beginning
        q |= Q(**{kw('istartswith'): f'{t},'})
        q |= Q(**{kw('istartswith'): f'{t} ,'})
        q |= Q(**{kw('istartswith'): f'NA\\{t},'})
        q |= Q(**{kw('istartswith'): f'NA/{t},'})
        # 3. Comma-separated: at the end
        q |= Q(**{kw('iendswith'): f',{t}'})
        q |= Q(**{kw('iendswith'): f', {t}'})
        q |= Q(**{kw('iendswith'): f',NA\\{t}'})
        q |= Q(**{kw('iendswith'): f', NA\\{t}'})
        q |= Q(**{kw('iendswith'): f',NA/{t}'})
        q |= Q(**{kw('iendswith'): f', NA/{t}'})
        # 4. Comma-separated: in the middle
        q |= Q(**{kw('icontains'): f',{t},'})
        q |= Q(**{kw('icontains'): f', {t},'})
        q |= Q(**{kw('icontains'): f', {t} ,'})
        q |= Q(**{kw('icontains'): f',NA\\{t},'})
        q |= Q(**{kw('icontains'): f', NA\\{t},'})
        q |= Q(**{kw('icontains'): f',NA/{t},'})
        q |= Q(**{kw('icontains'): f', NA/{t},'})
    return q

def matches_exact_location(emp_loc, allowed_locs):
    """
    In-memory check: returns True if any comma-separated token in emp_loc
    exactly matches (case-insensitive, ignoring NA\\ / NA/ prefix) any token in allowed_locs.
    """
    if not emp_loc or not allowed_locs:
        return False
    if isinstance(allowed_locs, str):
        allowed_locs = [allowed_locs]
    emp_tokens = set()
    for t in str(emp_loc).split(','):
        t_str = t.strip()
        if t_str:
            emp_tokens.add(t_str.lower())
            c = clean_loc_token(t_str)
            if c:
                emp_tokens.add(c.lower())
    for al in allowed_locs:
        al_raw = str(al).strip().lower()
        al_clean = clean_loc_token(al).lower()
        if al_raw in emp_tokens or al_clean in emp_tokens:
            return True
    return False


def parse_excel_date(raw_val):
    if not raw_val:
        return datetime.date.today()
    if isinstance(raw_val, datetime.datetime):
        return raw_val.date()
    if isinstance(raw_val, datetime.date):
        return raw_val
    if isinstance(raw_val, (int, float)):
        try:
            import openpyxl.utils.datetime as op_dt
            return op_dt.from_excel(raw_val).date()
        except Exception:
            try:
                return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(raw_val))
            except Exception:
                return datetime.date.today()
    if isinstance(raw_val, str):
        cleaned = raw_val.strip()
        if not cleaned:
            return datetime.date.today()
        formats = (
            '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%m/%d/%Y', '%Y/%m/%d',
            '%d-%b-%Y', '%d-%b-%y', '%d/%b/%Y', '%d/%b/%y',
            '%d-%B-%Y', '%d-%B-%y', '%d.%m.%Y', '%Y.%m.%d',
            '%d %b %Y', '%d %B %Y', '%d-%m-%y', '%d/%m/%y'
        )
        for fmt in formats:
            try:
                return datetime.datetime.strptime(cleaned, fmt).date()
            except ValueError:
                continue
    return datetime.date.today()


def get_column_indices(headers, col_mapping):
    indices = {}
    normalized_headers = [str(h or '').strip().lower().replace('_', ' ').replace('-', ' ').replace('.', '') for h in headers]
    for key, aliases in col_mapping.items():
        found = False
        for alias in aliases:
            norm_alias = alias.lower().replace('_', ' ').replace('-', ' ').replace('.', '')
            if norm_alias in normalized_headers:
                indices[key] = normalized_headers.index(norm_alias)
                found = True
                break
        if not found:
            for i, h in enumerate(normalized_headers):
                if any(a in h for a in aliases):
                    indices[key] = i
                    found = True
                    break
    return indices


class SimpleHTMLTableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.current_row = []
        self.current_cell = []
        self.in_cell = False

    def handle_starttag(self, tag, attrs):
        if tag.lower() in ('td', 'th'):
            self.in_cell = True
            self.current_cell = []

    def handle_endtag(self, tag):
        if tag.lower() in ('td', 'th'):
            self.in_cell = False
            self.current_row.append(''.join(self.current_cell).strip())
        elif tag.lower() == 'tr':
            if any(self.current_row):
                self.rows.append(self.current_row)
            self.current_row = []

    def handle_data(self, data):
        if self.in_cell:
            self.current_cell.append(data)


def extract_tabular_data(uploaded_file):
    """
    Universally extracts (headers, rows) from:
    1. .xlsx (Office Open XML via openpyxl)
    2. .xls (Excel 97-2004 binary BIFF via xlrd)
    3. HTML-table based .xls exports (via SimpleHTMLTableParser)
    4. XML Spreadsheet 2003 (SpreadsheetML)
    5. CSV / TSV text files
    Returns:
        headers: list of strings (row 1)
        rows: list of lists of values (rows 2+)
    """
    raw_content = uploaded_file.read()
    filename = getattr(uploaded_file, 'name', '').lower()

    # 1. If not explicitly .xls, try openpyxl first (.xlsx format)
    if not filename.endswith('.xls'):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw_content), read_only=True, data_only=True)
            sheet = wb.active
            all_rows = list(sheet.iter_rows(values_only=True))
            if all_rows:
                headers = [str(c).strip() if c is not None else "" for c in all_rows[0]]
                rows = all_rows[1:]
                return headers, rows
        except Exception:
            pass

    # 2. Try xlrd for .xls binary BIFF format
    try:
        book = xlrd.open_workbook(file_contents=raw_content)
        sheet = book.sheet_by_index(0)
        all_rows = []
        for r in range(sheet.nrows):
            row_vals = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        date_tuple = xlrd.xldate_as_tuple(cell.value, book.datemode)
                        val = datetime.date(*date_tuple[:3])
                    except Exception:
                        val = cell.value
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    if cell.value.is_integer():
                        val = int(cell.value)
                    else:
                        val = cell.value
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    val = bool(cell.value)
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    val = None
                else:
                    val = cell.value
                row_vals.append(val)
            all_rows.append(row_vals)
        if all_rows:
            headers = [str(c).strip() if c is not None else "" for c in all_rows[0]]
            rows = all_rows[1:]
            return headers, rows
    except Exception:
        pass

    # 3. HTML Table export (common when HR/ERP systems export HTML with .xls extension)
    sample_text = ""
    try:
        sample_text = raw_content[:4096].decode('utf-8', errors='ignore')
    except Exception:
        pass

    if '<table' in sample_text.lower() or '<html' in sample_text.lower():
        try:
            full_html = raw_content.decode('utf-8', errors='ignore')
            parser = SimpleHTMLTableParser()
            parser.feed(full_html)
            if parser.rows:
                headers = [str(c).strip() for c in parser.rows[0]]
                rows = parser.rows[1:]
                return headers, rows
        except Exception:
            pass

    # 4. XML Spreadsheet 2003 (SpreadsheetML)
    if 'urn:schemas-microsoft-com:office:spreadsheet' in sample_text:
        try:
            root = ET.fromstring(raw_content)
            all_rows = []
            for row_el in root.findall('.//{urn:schemas-microsoft-com:office:spreadsheet}Row'):
                row_vals = []
                for cell_el in row_el.findall('{urn:schemas-microsoft-com:office:spreadsheet}Cell'):
                    data_el = cell_el.find('{urn:schemas-microsoft-com:office:spreadsheet}Data')
                    row_vals.append(data_el.text if data_el is not None and data_el.text else '')
                if any(row_vals):
                    all_rows.append(row_vals)
            if all_rows:
                headers = [str(c).strip() for c in all_rows[0]]
                rows = all_rows[1:]
                return headers, rows
        except Exception:
            pass

    # 5. Fallback: try openpyxl in case file was named .xls but is actually OOXML / .xlsx
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_content), read_only=True, data_only=True)
        sheet = wb.active
        all_rows = list(sheet.iter_rows(values_only=True))
        if all_rows:
            headers = [str(c).strip() if c is not None else "" for c in all_rows[0]]
            rows = all_rows[1:]
            return headers, rows
    except Exception:
        pass

    # 6. CSV / TSV fallback
    try:
        text = raw_content.decode('utf-8', errors='ignore')
        reader = csv.reader(io.StringIO(text))
        all_rows = [r for r in reader if any(r)]
        if all_rows:
            headers = [str(c).strip() for c in all_rows[0]]
            rows = all_rows[1:]
            return headers, rows
    except Exception:
        pass

    raise ValueError("File is not a supported Excel or spreadsheet format. Please upload a valid .xlsx or .xls file.")


class BranchViewSet(viewsets.ModelViewSet):
    queryset = Branch.objects.all().order_by('name')
    serializer_class = BranchSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        queryset = super().get_queryset()
        user = self.request.user
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            dept_admin = user.department_admin
            if dept_admin.branches:
                queryset = queryset.filter(name__in=dept_admin.branches)
            else:
                queryset = queryset.none()
        return queryset


class DepartmentAdminViewSet(viewsets.ModelViewSet):
    queryset = DepartmentAdmin.objects.all().order_by('user__username')
    serializer_class = DepartmentAdminSerializer
    permission_classes = [permissions.IsAdminUser]


class CustomAuthToken(ObtainAuthToken):
    authentication_classes = ()
    permission_classes = ()

    def post(self, request, *args, **kwargs):
        serializer = self.serializer_class(data=request.data,
                                           context={'request': request})
        serializer.is_valid(raise_exception=True)
        user = serializer.validated_data['user']
        token, created = Token.objects.get_or_create(user=user)
        
        role = 'employee'
        departments = []
        locations = []
        branches = []
        if user.is_superuser:
            role = 'superuser'
            branches = list(Branch.objects.values_list('name', flat=True))
        elif hasattr(user, 'department_admin'):
            role = 'department_admin'
            departments = user.department_admin.departments or []
            locations = user.department_admin.locations or []
            branches = user.department_admin.branches or []

        return Response({
            'token': token.key,
            'user_id': user.pk,
            'username': user.username,
            'is_staff': user.is_staff or user.is_superuser or hasattr(user, 'department_admin'),
            'role': role,
            'departments': departments,
            'locations': locations,
            'branches': branches,
            # Backwards compatibility fallback (first department or None)
            'department': departments[0] if departments else None
        })


class CurrentUserView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        user = request.user
        role = 'employee'
        departments = []
        locations = []
        branches = []
        if user.is_superuser:
            role = 'superuser'
            branches = list(Branch.objects.values_list('name', flat=True))
        elif hasattr(user, 'department_admin'):
            role = 'department_admin'
            departments = user.department_admin.departments or []
            locations = user.department_admin.locations or []
            branches = user.department_admin.branches or []

        return Response({
            'username': user.username,
            'is_staff': user.is_staff or user.is_superuser or hasattr(user, 'department_admin'),
            'role': role,
            'departments': departments,
            'locations': locations,
            'branches': branches,
            'department': departments[0] if departments else None
        })



class PasswordResetView(APIView):
    authentication_classes = ()
    permission_classes = [permissions.AllowAny]

    def post(self, request, *args, **kwargs):
        username = request.data.get('username')
        current_password = request.data.get('current_password')
        new_password = request.data.get('new_password')

        if not username or not current_password or not new_password:
            return Response(
                {'error': 'Username, current password, and new password are required.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            user = User.objects.get(username=username)
        except User.DoesNotExist:
            return Response(
                {'error': 'User not found.'},
                status=status.HTTP_404_NOT_FOUND
            )

        # Verify current password
        if not user.check_password(current_password):
            return Response(
                {'error': 'Incorrect current password.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # If verification passed, reset password
        user.set_password(new_password)
        user.save()
        
        # Delete existing auth token if exists to force relogin
        Token.objects.filter(user=user).delete()

        return Response(
            {'message': 'Password has been reset successfully.'},
            status=status.HTTP_200_OK
        )


class EmployeeViewSet(viewsets.ModelViewSet):
    queryset = Employee.objects.all().order_by('employee_code')
    serializer_class = EmployeeSerializer

    def get_permissions(self):
        # Allow any user to retrieve or list (needed for public appraisal form dropdowns)
        if self.action in ['list', 'retrieve']:
            return [permissions.AllowAny()]
        return [IsAdminUserOrReadOnly()]

    def create(self, request, *args, **kwargs):
        emp_code = str(request.data.get('employee_code') or '').strip()
        name = str(request.data.get('name') or '').strip()
        dept = str(request.data.get('department') or '').strip()
        loc = str(request.data.get('location') or '').strip()

        if not emp_code:
            return Response({"employee_code": ["Employee Code is required."]}, status=status.HTTP_400_BAD_REQUEST)
        if not name:
            return Response({"name": ["Employee Name is required."]}, status=status.HTTP_400_BAD_REQUEST)
        if not dept:
            return Response({"department": ["Department is required."]}, status=status.HTTP_400_BAD_REQUEST)

        # Department admin permission check (case-insensitive)
        user = request.user
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            allowed_depts = [d.strip().lower() for d in user.department_admin.departments]
            req_depts = [d.strip().lower() for d in dept.split(',') if d.strip()]
            if not any(rd in allowed_depts for rd in req_depts):
                return Response({"department": [f"You do not have permission to add employees in department '{dept}'."]}, status=status.HTTP_403_FORBIDDEN)

        # Upsert: If employee with (employee_code, department, location) exists, update it!
        existing = Employee.objects.filter(employee_code=emp_code, department=dept, location=loc).first()
        if not existing and Employee.objects.filter(employee_code=emp_code).count() == 1:
            existing = Employee.objects.filter(employee_code=emp_code).first()

        if existing:
            serializer = self.get_serializer(existing, data=request.data, partial=True)
        else:
            serializer = self.get_serializer(data=request.data)

        serializer.is_valid(raise_exception=True)
        emp = serializer.save()

        # Ensure user login account exists with default password Iqraa@123
        if not User.objects.filter(username=emp.employee_code).exists():
            User.objects.create_user(username=emp.employee_code, password='Iqraa@123', is_staff=False)

        headers = self.get_success_headers(serializer.data)
        return Response(serializer.data, status=status.HTTP_201_CREATED if not existing else status.HTTP_200_OK, headers=headers)


    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        queryset = self.filter_queryset(self.get_queryset())
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Employee Master"

        headers = ["Employee Code", "Employee Name", "Joining Date", "Organization", "Division", "Category", "Department", "Designation", "Duty Location", "Branch", "Level", "Assignment Period", "Status"]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

        for row_num, obj in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=obj.employee_code)
            ws.cell(row=row_num, column=2, value=obj.name)
            ws.cell(row=row_num, column=3, value=obj.date_of_joining.strftime('%Y-%m-%d') if obj.date_of_joining else '')
            ws.cell(row=row_num, column=4, value=obj.organization or '')
            ws.cell(row=row_num, column=5, value=obj.division or '')
            ws.cell(row=row_num, column=6, value=obj.category or '')
            ws.cell(row=row_num, column=7, value=obj.department)
            ws.cell(row=row_num, column=8, value=obj.designation)
            ws.cell(row=row_num, column=9, value=obj.location)
            ws.cell(row=row_num, column=10, value=obj.branch.name if obj.branch else '')
            ws.cell(row=row_num, column=11, value=obj.level or '')
            ws.cell(row=row_num, column=12, value=obj.assignment_period)
            ws.cell(row=row_num, column=13, value=obj.status)

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 12)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=employees_master.xlsx"
        wb.save(response)
        return response

    def get_queryset(self):
        queryset = super().get_queryset()
        user = self.request.user

        if user and user.is_authenticated and hasattr(user, 'department_admin'):
            queryset = queryset.exclude(employee_code=user.username)
            dept_admin = user.department_admin

            # Strict white-list: if any of department, location, or branch is untagged, grant no access
            if not dept_admin.departments or not dept_admin.locations or not dept_admin.branches:
                return queryset.none()

            dept_filter = Q()
            for d in dept_admin.departments:
                dept_filter |= Q(department__iexact=d) | Q(department__icontains=d)
            queryset = queryset.filter(dept_filter)

            loc_filter = Q()
            for loc in dept_admin.locations:
                loc_filter |= build_exact_location_filter(loc)
            queryset = queryset.filter(loc_filter)

            br_filter = Q()
            for br in dept_admin.branches:
                br_filter |= Q(branch__name__iexact=br)
            queryset = queryset.filter(br_filter)

            # Restrict department admin to ONLY visible staff in the active period
            active_p = IncentivePeriod.objects.filter(is_active=True).first()
            if active_p:
                queryset = queryset.filter(
                    Q(incentive_period=active_p) | 
                    Q(assignment_period__iexact=active_p.name) | 
                    Q(assignment_period__iexact=active_p.months)
                )

        dept = self.request.query_params.get('department')
        location_param = self.request.query_params.get('location')
        branch_param = self.request.query_params.get('branch')
        desig = self.request.query_params.get('designation')
        status_param = self.request.query_params.get('status')
        search = self.request.query_params.get('search')
        period_param = self.request.query_params.get('period') or self.request.query_params.get('period_id') or self.request.query_params.get('incentive_period') or self.request.query_params.get('assignment_period')

        if period_param and str(period_param).strip().lower() != 'all':
            if str(period_param).strip().lower() == 'active':
                active_p = IncentivePeriod.objects.filter(is_active=True).first()
                if active_p:
                    queryset = queryset.filter(
                        Q(incentive_period=active_p) | 
                        Q(assignment_period__iexact=active_p.name) | 
                        Q(assignment_period__iexact=active_p.months)
                    )
            elif str(period_param).isdigit():
                p_obj = IncentivePeriod.objects.filter(id=int(period_param)).first()
                if p_obj:
                    queryset = queryset.filter(
                        Q(incentive_period=p_obj) | 
                        Q(assignment_period__iexact=p_obj.name) | 
                        Q(assignment_period__iexact=p_obj.months)
                    )
                else:
                    queryset = queryset.filter(incentive_period_id=int(period_param))
            else:
                queryset = queryset.filter(
                    Q(assignment_period__iexact=period_param) | 
                    Q(incentive_period__name__iexact=period_param)
                )

        trashed_param = self.request.query_params.get('trashed') or self.request.query_params.get('is_trashed')
        if getattr(self, 'action', None) in ['restore', 'permanent_delete']:
            pass
        elif str(trashed_param).strip().lower() in ['true', '1', 'yes']:
            queryset = queryset.filter(is_trashed=True)
        elif str(trashed_param).strip().lower() == 'all':
            pass
        else:
            queryset = queryset.filter(is_trashed=False)

        if dept:
            queryset = queryset.filter(Q(department__iexact=dept) | Q(department__icontains=dept))
        if location_param:
            queryset = queryset.filter(build_exact_location_filter(location_param))
        if branch_param:
            queryset = queryset.filter(branch__name__iexact=branch_param)
        if desig:
            queryset = queryset.filter(designation__iexact=desig)
        if status_param:
            queryset = queryset.filter(status=status_param)
        if search:
            clean_search = str(search).strip()
            if clean_search:
                search_q = (
                    Q(employee_code__iexact=clean_search) | 
                    Q(employee_code__icontains=clean_search) | 
                    Q(name__icontains=clean_search)
                )
                if clean_search.isdigit():
                    search_q |= Q(employee_code__iexact=clean_search.zfill(6))
                    search_q |= Q(employee_code__iexact=clean_search.zfill(5))
                    search_q |= Q(employee_code__iexact=clean_search.zfill(4))
                    search_q |= Q(employee_code__endswith=clean_search)
                queryset = queryset.filter(search_q)
        return queryset

    def perform_destroy(self, instance):
        # Soft-delete to trash: preserve in DB while removing from active frontend
        instance.is_trashed = True
        instance.trashed_at = timezone.now()
        instance.status = 'Trashed'
        instance.save(update_fields=['is_trashed', 'trashed_at', 'status'])

    @action(detail=True, methods=['post'], url_path='restore')
    def restore(self, request, pk=None):
        instance = Employee.objects.filter(pk=pk).first()
        if not instance:
            return Response({"error": "Employee not found"}, status=status.HTTP_404_NOT_FOUND)
        user = request.user
        if user and user.is_authenticated and hasattr(user, 'department_admin'):
            if instance.department not in user.department_admin.departments:
                return Response({"error": "Permission denied"}, status=status.HTTP_403_FORBIDDEN)
        instance.is_trashed = False
        instance.trashed_at = None
        instance.status = 'Active'
        instance.save(update_fields=['is_trashed', 'trashed_at', 'status'])
        return Response({
            "message": f"Employee '{instance.employee_code}' restored to active staff.",
            "employee": self.get_serializer(instance).data
        }, status=status.HTTP_200_OK)

    @action(detail=True, methods=['delete', 'post'], url_path='permanent-delete')
    def permanent_delete(self, request, pk=None):
        instance = Employee.objects.filter(pk=pk).first()
        if not instance:
            return Response({"error": "Employee not found"}, status=status.HTTP_404_NOT_FOUND)
        user = request.user
        if user and user.is_authenticated and hasattr(user, 'department_admin'):
            if instance.department not in user.department_admin.departments:
                return Response({"error": "Permission denied"}, status=status.HTTP_403_FORBIDDEN)
        code = instance.employee_code
        instance.delete()
        return Response({
            "message": f"Employee '{code}' permanently deleted from database."
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['post'], url_path='bulk-trash')
    def bulk_trash(self, request):
        employee_ids = request.data.get('employee_ids')
        period_id = request.data.get('period_id')
        branch_name = request.data.get('branch')

        qs = Employee.objects.filter(is_trashed=False)
        user = request.user
        if user and user.is_authenticated and hasattr(user, 'department_admin'):
            qs = qs.filter(department__in=user.department_admin.departments)

        if employee_ids and isinstance(employee_ids, list):
            qs = qs.filter(id__in=employee_ids)
        elif period_id and str(period_id).strip().lower() != 'all':
            if str(period_id).isdigit():
                qs = qs.filter(Q(incentive_period_id=int(period_id)) | Q(assignment_period=str(period_id)))
            else:
                qs = qs.filter(Q(assignment_period__iexact=str(period_id)) | Q(incentive_period__name__iexact=str(period_id)))
        elif branch_name:
            qs = qs.filter(branch__name=branch_name)

        now = timezone.now()
        count = qs.update(is_trashed=True, trashed_at=now, status='Trashed')
        return Response({
            "message": f"Successfully moved {count} employee(s) to trash.",
            "count": count
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['post'], url_path='restore-all')
    def restore_all(self, request):
        qs = Employee.objects.filter(is_trashed=True)
        user = request.user
        if user and user.is_authenticated and hasattr(user, 'department_admin'):
            qs = qs.filter(department__in=user.department_admin.departments)
        count = qs.update(is_trashed=False, trashed_at=None, status='Active')
        return Response({
            "message": f"Successfully restored {count} employee(s) to active directory.",
            "count": count
        }, status=status.HTTP_200_OK)

    @action(detail=False, methods=['post', 'delete'], url_path='empty-trash')
    def empty_trash(self, request):
        qs = Employee.objects.filter(is_trashed=True)
        user = request.user
        if user and user.is_authenticated and hasattr(user, 'department_admin'):
            qs = qs.filter(department__in=user.department_admin.departments)
        count = qs.count()
        qs.delete()
        return Response({
            "message": f"Permanently deleted {count} trashed employee(s).",
            "count": count
        }, status=status.HTTP_200_OK)

    def perform_create(self, serializer):
        user = self.request.user
        if user and user.is_authenticated and hasattr(user, 'department_admin'):
            dept = serializer.validated_data.get('department')
            loc = serializer.validated_data.get('location') or ""
            branch_obj = serializer.validated_data.get('branch')
            dept_admin = user.department_admin
            save_kwargs = {}
            
            allowed_depts = dept_admin.departments
            if allowed_depts and dept not in allowed_depts:
                save_kwargs['department'] = allowed_depts[0]
                
            allowed_locs = dept_admin.locations
            if allowed_locs:
                loc_is_allowed = any(al.lower() in loc.lower() for al in allowed_locs)
                if not loc_is_allowed:
                    save_kwargs['location'] = allowed_locs[0]

            allowed_branches = dept_admin.branches
            if allowed_branches:
                br_name = branch_obj.name if branch_obj else ""
                br_is_allowed = any(ab.lower() in br_name.lower() for ab in allowed_branches)
                if not br_is_allowed:
                    fallback_br, _ = Branch.objects.get_or_create(name=allowed_branches[0])
                    save_kwargs['branch'] = fallback_br
                
            serializer.save(**save_kwargs)
        else:
            serializer.save()

    def perform_update(self, serializer):
        user = self.request.user
        if user and user.is_authenticated and hasattr(user, 'department_admin'):
            dept = serializer.validated_data.get('department')
            loc = serializer.validated_data.get('location') or ""
            branch_obj = serializer.validated_data.get('branch')
            dept_admin = user.department_admin
            save_kwargs = {}
            
            allowed_depts = dept_admin.departments
            if allowed_depts and dept not in allowed_depts:
                save_kwargs['department'] = allowed_depts[0]
                
            allowed_locs = dept_admin.locations
            if allowed_locs:
                loc_is_allowed = any(al.lower() in loc.lower() for al in allowed_locs)
                if not loc_is_allowed:
                    save_kwargs['location'] = allowed_locs[0]

            allowed_branches = dept_admin.branches
            if allowed_branches:
                br_name = branch_obj.name if branch_obj else ""
                br_is_allowed = any(ab.lower() in br_name.lower() for ab in allowed_branches)
                if not br_is_allowed:
                    fallback_br, _ = Branch.objects.get_or_create(name=allowed_branches[0])
                    save_kwargs['branch'] = fallback_br
                
            serializer.save(**save_kwargs)
        else:
            serializer.save()


class EmployeeAutoFetchView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        code = request.query_params.get('code')
        dept_param = request.query_params.get('department')
        loc_param = request.query_params.get('location')
        branch_param = request.query_params.get('branch')
        if not code:
            return Response({"error": "Employee code is required"}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            qs = Employee.objects.filter(employee_code=code, is_trashed=False)
            if dept_param:
                qs = qs.filter(Q(department__iexact=dept_param) | Q(department__icontains=dept_param))
            if loc_param:
                qs = qs.filter(build_exact_location_filter(loc_param))
            if branch_param:
                qs = qs.filter(branch__name__iexact=branch_param)
            
            if not qs.exists():
                raise Employee.DoesNotExist

            # Enforce department, location, and branch restriction for department admins
            user = request.user
            if user and user.is_authenticated and hasattr(user, 'department_admin'):
                if code.strip().lower() == user.username.strip().lower():
                    return Response({"error": "Self-appraisal is not permitted. Department administrators cannot appraise their own assessment."}, status=status.HTTP_403_FORBIDDEN)
                dept_admin = user.department_admin

                if not dept_admin.departments:
                    return Response({"error": "No authorized departments assigned to your administrator account."}, status=status.HTTP_403_FORBIDDEN)
                
                # Check if there is any record within the admin's departments
                dept_filter_q = Q()
                for d in dept_admin.departments:
                    dept_filter_q |= Q(department__iexact=d) | Q(department__icontains=d)
                allowed_qs = qs.filter(dept_filter_q)
                if not allowed_qs.exists():
                    return Response({"error": "You do not have permission to access employees outside your department."}, status=status.HTTP_403_FORBIDDEN)
                
                if not dept_admin.locations:
                    return Response({"error": "No authorized duty locations assigned to your administrator account."}, status=status.HTTP_403_FORBIDDEN)

                loc_filtered = []
                for emp in allowed_qs:
                    if matches_exact_location(emp.location, dept_admin.locations):
                        loc_filtered.append(emp)
                
                if not loc_filtered:
                    return Response({"error": "You do not have permission to access employees outside your assigned duty locations."}, status=status.HTTP_403_FORBIDDEN)
                allowed_qs = Employee.objects.filter(id__in=[e.id for e in loc_filtered])

                if not dept_admin.branches:
                    return Response({"error": "No authorized hospital branches assigned to your administrator account."}, status=status.HTTP_403_FORBIDDEN)

                br_filtered = []
                for emp in allowed_qs:
                    emp_br = (emp.branch.name if emp.branch else "").lower()
                    if any(ab.lower() in emp_br for ab in dept_admin.branches):
                        br_filtered.append(emp)

                if not br_filtered:
                    return Response({"error": "You do not have permission to access employees outside your assigned branches."}, status=status.HTTP_403_FORBIDDEN)
                allowed_qs = Employee.objects.filter(id__in=[e.id for e in br_filtered])

                active_allowed = allowed_qs.filter(status='Active')
                employee = active_allowed.first() or allowed_qs.first()
            else:
                # Not a department admin (e.g. superuser), just find an active one if possible
                active_qs = qs.filter(status='Active')
                employee = active_qs.first() or qs.first()

            # Check if active
            if employee.status != 'Active':
                return Response({"error": "Employee is inactive"}, status=status.HTTP_400_BAD_REQUEST)

            # For department admin, ensure the employee belongs to the active period
            if user and user.is_authenticated and hasattr(user, 'department_admin'):
                active_p = IncentivePeriod.objects.filter(is_active=True).first()
                if active_p:
                    period_match = (
                        employee.incentive_period_id == active_p.id or
                        (employee.assignment_period and (
                            employee.assignment_period.strip().lower() == active_p.name.strip().lower() or
                            employee.assignment_period.strip().lower() == (active_p.months or '').strip().lower()
                        ))
                    )
                    if not period_match:
                        return Response({
                            "error": f"Employee '{code}' is not assigned to the active appraisal cycle ({active_p.name})."
                        }, status=status.HTTP_400_BAD_REQUEST)

            serializer = EmployeeSerializer(employee)
            department = employee.department
            location = employee.location
            branch_obj = employee.branch
            
            # Active staff counts (excluding the logged-in department admin evaluating the team)
            target_department = dept_param if (dept_param and dept_param.lower() in (department or '').lower()) else department
            target_department_primary = target_department.split(',')[0].strip() if target_department else ''
            dept_active_qs = Employee.objects.filter(Q(department__iexact=target_department_primary) | Q(department__icontains=target_department_primary), status='Active') if target_department_primary else Employee.objects.none()
            if branch_obj:
                dept_active_qs = dept_active_qs.filter(branch=branch_obj)
            if user and user.is_authenticated and hasattr(user, 'department_admin'):
                dept_active_qs = dept_active_qs.exclude(employee_code=user.username)
            dept_total_active = dept_active_qs.count()

            # Location active staff scoped to this department and location (excluding the logged-in department admin)
            target_location = loc_param if (loc_param and matches_exact_location(location, [loc_param])) else location
            loc_filter = {'status': 'Active'}
            if department:
                loc_filter['department'] = department
            loc_active_qs = Employee.objects.filter(**loc_filter) if target_location else Employee.objects.none()
            if target_location:
                loc_active_qs = loc_active_qs.filter(build_exact_location_filter(target_location))
            if branch_obj:
                loc_active_qs = loc_active_qs.filter(branch=branch_obj)
            if user and user.is_authenticated and hasattr(user, 'department_admin'):
                loc_active_qs = loc_active_qs.exclude(employee_code=user.username)
            loc_total_active = loc_active_qs.count()

            branch_total_active_qs = Employee.objects.filter(branch=branch_obj, status='Active') if branch_obj else Employee.objects.none()
            if user and user.is_authenticated and hasattr(user, 'department_admin'):
                branch_total_active_qs = branch_total_active_qs.exclude(employee_code=user.username)
            branch_total_active = branch_total_active_qs.count()

            # Determine primary scope for limits (location takes precedence if employee has a location with active staff in this department, else department)
            use_location = bool(location and loc_total_active > 0)
            scope_staff = loc_total_active if use_location else dept_total_active

            # Rule calculation:
            # 1. Staff <= 10: strictly max min(5, ceil(N/2)) staff for score 2.70 - 3.00 (90% - 100%), others < 2.70
            # 2. Staff > 10: max 50% (half) of active staff for score 2.25 - 3.00 (75% - 100%), others < 2.25
            if scope_staff <= 10:
                rule_type = 'lte_10'
                threshold_score = 2.70
                threshold_marks = 90
                max_allowed = min(5, (scope_staff + 1) // 2) if scope_staff > 0 else 1
                threshold_label = '2.7 - 3.0 (90% - 100%)'
            else:
                rule_type = 'gt_10'
                threshold_score = 2.25
                threshold_marks = 75
                max_allowed = max(1, (scope_staff + 1) // 2)
                threshold_label = '2.25 - 3.0 (75% - 100%)'

            active_p = IncentivePeriod.objects.filter(is_active=True).first()

            def count_restricted(scope_field, scope_val, thr_score, thr_marks):
                if not scope_val:
                    return 0
                filter_kwargs = {scope_field: scope_val}
                if scope_field == 'location' and department:
                    filter_kwargs['department'] = department
                q = Appraisal.objects.filter(**filter_kwargs)
                if user and user.is_authenticated and hasattr(user, 'department_admin'):
                    q = q.exclude(employee_code=user.username)
                if branch_obj:
                    q = q.filter(branch=branch_obj)
                if active_p:
                    period_filter = Q(assignment_period__iexact=active_p.name)
                    if active_p.months:
                        period_filter |= Q(assignment_period__iexact=active_p.months)
                    if employee.assignment_period and employee.assignment_period.strip().lower() not in [active_p.name.strip().lower(), (active_p.months or '').strip().lower()]:
                        period_filter |= Q(assignment_period__iexact=employee.assignment_period)
                    q = q.filter(period_filter)
                elif employee.assignment_period:
                    q = q.filter(assignment_period__iexact=employee.assignment_period)

                if thr_score == 2.70:
                    return q.filter(Q(final_score__gte=2.70) | Q(marks_obtained__gte=90) | Q(rating__icontains='Outstanding')).count()
                else:
                    return q.filter(Q(final_score__gte=2.25) | Q(marks_obtained__gte=75) | Q(rating__icontains='Outstanding') | Q(rating__icontains='Very Good')).count()

            dept_existing = count_restricted('department', department, threshold_score, threshold_marks)
            loc_existing = count_restricted('location', location, threshold_score, threshold_marks) if location else 0

            data = serializer.data
            if dept_param and (dept_param.lower() in (department or '').lower()):
                data['department'] = dept_param
            if loc_param and (loc_param.lower() in (location or '').lower()):
                data['location'] = loc_param
            if active_p and (not data.get('assignment_period') or data.get('assignment_period') in ['Annual', 'Monthly', 'Quarterly', 'Semi-Annual']):
                data['assignment_period'] = active_p.name
            data['active_period_name'] = active_p.name if active_p else None
            data['department_total_active'] = dept_total_active
            data['department_existing_ab'] = dept_existing
            data['location_total_active'] = loc_total_active
            data['location_existing_ab'] = loc_existing
            data['branch_total_active'] = branch_total_active
            data['rule_type'] = rule_type
            data['threshold_score'] = threshold_score
            data['threshold_marks'] = threshold_marks
            data['threshold_label'] = threshold_label
            data['max_allowed'] = max_allowed
            data['scope_total_active'] = scope_staff
            data['scope_existing_restricted'] = loc_existing if use_location else dept_existing
            return Response(data)
        except Employee.DoesNotExist:
            return Response({"error": f"Employee with code '{code}' not found in master database."}, status=status.HTTP_404_NOT_FOUND)


class ExcelImportView(APIView):
    permission_classes = [IsAdminUserOrReadOnly]

    def post(self, request):
        if 'file' not in request.FILES:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)
        
        excel_file = request.FILES['file']
        try:
            headers, rows = extract_tabular_data(excel_file)
        except Exception as e:
            return Response({"error": f"Invalid excel file format: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)

        # Map expected headers to indexes (broad matching)
        col_mapping = {
            'employee_code': ['employee code', 'code', 'emp code', 'emp_code', 'employee_code', 'emp id', 'empid', 'emp_id', 'employee id', 'staff id', 'staff code', 'staff no', 'trainee code', 'token no', 'id'],
            'name': ['employee name', 'name', 'emp name', 'emp_name', 'employee_name', 'staff name', 'trainee name', 'full name', 'staff'],
            'date_of_joining': ['joining date', 'date of joining', 'doj', 'joining_date', 'date_of_joining', 'join date', 'date of join', 'd.o.j', 'training start date', 'start date'],
            'branch': ['branch', 'branch name', 'branches', 'hospital branch', 'duty branch', 'hospital unit', 'unit', 'brach', 'hospital'],
            'organization': ['organization', 'org', 'company', 'institution'],
            'division': ['division', 'div', 'division name'],
            'category': ['category', 'cat', 'staff category', 'emp category'],
            'department': ['department', 'dept', 'department name', 'dept name', 'dept.', 'department_name'],
            'designation': ['designation', 'desig', 'role', 'post', 'position', 'job title', 'title'],
            'location': ['duty location', 'location', 'loc', 'duty station', 'work location', 'location name', 'ward', 'area', 'place', 'station'],
            'level': ['level', 'emp level', 'employee level', 'staff level', 'grade level']
        }

        indices = get_column_indices(headers, col_mapping)
        required_keys = ['employee_code', 'name', 'department']
        for rk in required_keys:
            if rk not in indices:
                return Response({
                    "error": f"Missing required column matching: {rk.replace('_', ' ').title()}. Please check Excel headers."
                }, status=status.HTTP_400_BAD_REQUEST)

        # Selected active branch fallback passed from frontend
        selected_branch = request.data.get('selected_branch') or request.query_params.get('branch') or 'IQRAA INTERNATIONAL HOSPITAL & RESEARCH CENTRE'

        # Target incentive period from request data or query param, or active period
        period_param = request.data.get('period_id') or request.query_params.get('period_id') or request.data.get('period')
        target_period = None
        if period_param:
            if str(period_param).isdigit():
                target_period = IncentivePeriod.objects.filter(id=int(period_param)).first()
            else:
                target_period = IncentivePeriod.objects.filter(name__iexact=str(period_param)).first()
        if not target_period:
            target_period = IncentivePeriod.objects.filter(is_active=True).first()

        period_name_str = target_period.name if target_period else 'Annual'

        # Process rows
        success_count = 0
        skipped_count = 0
        failed_rows = []
        imported_meta = {} # dept -> {'locations': set(), 'branches': set()}
        
        row_num = 1
        for row_data in rows:
            row_num += 1
            # Skip completely empty rows
            if not any(row_data):
                continue
            
            try:
                raw_code = row_data[indices['employee_code']] if indices.get('employee_code') is not None and indices['employee_code'] < len(row_data) else ""
                if isinstance(raw_code, float) and raw_code.is_integer():
                    emp_code = str(int(raw_code)).strip()
                else:
                    emp_code = str(raw_code or "").strip()
                    if emp_code.endswith('.0') and emp_code[:-2].replace('-', '').isdigit():
                        emp_code = emp_code[:-2]

                name = str(row_data[indices['name']]).strip() if indices.get('name') is not None and indices['name'] < len(row_data) and row_data[indices['name']] is not None else ""
                dept = str(row_data[indices['department']]).strip() if indices.get('department') is not None and indices['department'] < len(row_data) and row_data[indices['department']] is not None else ""
                location = str(row_data[indices['location']]).strip() if 'location' in indices and indices['location'] < len(row_data) and row_data[indices['location']] is not None else ""
                branch_str = str(row_data[indices['branch']]).strip() if 'branch' in indices and indices['branch'] < len(row_data) and row_data[indices['branch']] is not None else ""
                desig = str(row_data[indices['designation']]).strip() if 'designation' in indices and indices['designation'] < len(row_data) and row_data[indices['designation']] is not None else "Staff"
                org_str = str(row_data[indices['organization']]).strip() if 'organization' in indices and indices['organization'] < len(row_data) and row_data[indices['organization']] else ""
                div_str = str(row_data[indices['division']]).strip() if 'division' in indices and indices['division'] < len(row_data) and row_data[indices['division']] else ""
                cat_str = str(row_data[indices['category']]).strip() if 'category' in indices and indices['category'] < len(row_data) and row_data[indices['category']] else ""
                lvl_str = str(row_data[indices['level']]).strip() if 'level' in indices and indices['level'] < len(row_data) and row_data[indices['level']] else ""
                doj_raw = row_data[indices['date_of_joining']] if 'date_of_joining' in indices and indices['date_of_joining'] < len(row_data) else None
            except IndexError:
                failed_rows.append({
                    "row": row_num,
                    "code": "",
                    "reason": "Missing cell values matching header layout"
                })
                continue

            if not emp_code or not name:
                failed_rows.append({
                    "row": row_num,
                    "code": emp_code,
                    "reason": "Employee Code and Name cannot be empty"
                })
                continue

            # Parse Date of Joining
            doj = parse_excel_date(doj_raw)

            # If branch is not explicitly provided, check if organization matches an existing Branch
            if not branch_str and org_str:
                matched_br = Branch.objects.filter(name__iexact=org_str).first()
                if matched_br:
                    branch_str = matched_br.name

            # Fallback to selected_branch if branch_str is still empty
            if not branch_str:
                branch_str = selected_branch

            # Look up or create branch object
            branch_obj = None
            if branch_str:
                branch_obj, _ = Branch.objects.get_or_create(name=branch_str)

            # Auto-register unique division into DivisionRate if not exists
            if div_str:
                DivisionRate.objects.get_or_create(division_name=div_str, defaults={'rate': 0.0})

            # Auto-register unique employee level into EmployeeLevelConfig if not exists
            if lvl_str:
                EmployeeLevelConfig.objects.get_or_create(level_code=lvl_str, defaults={'description': f'Level {lvl_str}', 'score': 1.0})

            try:
                # Upsert employee (updates if exists, creates if new)
                emp_obj, created = Employee.objects.update_or_create(
                    employee_code=emp_code,
                    department=dept,
                    location=location,
                    defaults={
                        'name': name,
                        'organization': org_str,
                        'division': div_str,
                        'category': cat_str,
                        'level': lvl_str,
                        'branch': branch_obj,
                        'designation': desig or 'Staff',
                        'date_of_joining': doj,
                        'assignment_period': period_name_str,
                        'incentive_period': target_period,
                        'status': 'Active',
                        'is_trashed': False,
                        'trashed_at': None
                    }
                )
                success_count += 1
                
                # Auto-create user login credentials for the employee with default password Iqraa@123
                if not User.objects.filter(username=emp_code).exists():
                    User.objects.create_user(username=emp_code, password='Iqraa@123', is_staff=False)
            except Exception as ex:
                failed_rows.append({
                    "row": row_num,
                    "code": emp_code,
                    "reason": f"Database write error: {str(ex)}"
                })

        summary = {
            "total_processed": len(rows),
            "successfully_imported": success_count,
            "skipped_duplicates": skipped_count,
            "failed": len(failed_rows)
        }

        return Response({
            "summary": summary,
            "failed_rows": failed_rows
        }, status=status.HTTP_200_OK if len(failed_rows) == 0 else status.HTTP_207_MULTI_STATUS)



def get_pending_appraisal_records(user, params):
    if hasattr(user, 'department_admin'):
        dept_admin = user.department_admin
        if not dept_admin.departments or not dept_admin.locations or not dept_admin.branches:
            return []
        dept_filter = Q()
        for d in dept_admin.departments:
            dept_filter |= Q(department__iexact=d) | Q(department__icontains=d)
        emp_qs = Employee.objects.filter(dept_filter, status='Active', is_trashed=False)
        if user and user.is_authenticated:
            emp_qs = emp_qs.exclude(employee_code=user.username)
        loc_filter = Q()
        for loc in dept_admin.locations:
            loc_filter |= build_exact_location_filter(loc)
        emp_qs = emp_qs.filter(loc_filter)
        br_filter = Q()
        for br in dept_admin.branches:
            br_filter |= Q(branch__name__iexact=br)
        emp_qs = emp_qs.filter(br_filter)

        # Restrict department admin to ONLY visible staff in the active period
        active_p = IncentivePeriod.objects.filter(is_active=True).first()
        if active_p:
            emp_qs = emp_qs.filter(
                Q(incentive_period=active_p) | 
                Q(assignment_period__iexact=active_p.name) | 
                Q(assignment_period__iexact=active_p.months)
            )
    else:
        emp_qs = Employee.objects.filter(status='Active', is_trashed=False)

    app_submitted = Appraisal.objects.all()
    active_p = IncentivePeriod.objects.filter(is_active=True).first()
    target_period = params.get('assignment_period')
    if target_period:
        app_submitted = app_submitted.filter(assignment_period=target_period)
    elif hasattr(user, 'department_admin') and active_p:
        app_submitted = app_submitted.filter(
            Q(assignment_period__iexact=active_p.name) | 
            Q(assignment_period__iexact=active_p.months)
        )
    if hasattr(user, 'department_admin'):
        app_submitted = app_submitted.filter(dept_filter)
        if dept_admin.locations:
            app_submitted = app_submitted.filter(loc_filter)
        if dept_admin.branches:
            app_submitted = app_submitted.filter(br_filter)

    submitted_tuples = list(app_submitted.values_list('employee_code', 'department', 'location', 'branch__name'))
    
    # Exclude employee records that already have a submitted appraisal in that specific department, location & branch
    pending_emp_ids = []
    for emp in emp_qs:
        emp_br_name = emp.branch.name if emp.branch else ''
        emp_depts = [d.strip().lower() for d in (emp.department or '').split(',') if d.strip()]
        emp_locs = [l.strip().lower() for l in (emp.location or '').split(',') if l.strip()]

        is_submitted = False
        for sub_code, sub_dept, sub_loc, sub_br in submitted_tuples:
            if sub_code == emp.employee_code:
                if not emp_br_name or not sub_br or emp_br_name.lower() == sub_br.lower():
                    sub_dept_clean = (sub_dept or '').strip().lower()
                    sub_loc_clean = (sub_loc or '').strip().lower()
                    dept_match = not sub_dept_clean or (sub_dept_clean in emp_depts) or any(d in sub_dept_clean for d in emp_depts)
                    loc_match = not sub_loc_clean or matches_exact_location(emp.location, [sub_loc])
                    if dept_match and loc_match:
                        is_submitted = True
                        break
        if not is_submitted:
            pending_emp_ids.append(emp.id)

    emp_qs = emp_qs.filter(id__in=pending_emp_ids)

    dept = params.get('department')
    location_param = params.get('location')
    branch_param = params.get('branch')
    desig = params.get('designation')
    period = params.get('assignment_period')
    name = params.get('name')
    code = params.get('code')
    employee_code = params.get('employee_code')
    search = params.get('search')

    if dept:
        emp_qs = emp_qs.filter(Q(department__iexact=dept) | Q(department__icontains=dept))
    if location_param:
        emp_qs = emp_qs.filter(build_exact_location_filter(location_param))
    if branch_param:
        emp_qs = emp_qs.filter(branch__name__iexact=branch_param)
    if desig:
        emp_qs = emp_qs.filter(designation__iexact=desig)
    if period:
        emp_qs = emp_qs.filter(assignment_period=period)
    if name:
        emp_qs = emp_qs.filter(name__icontains=name)
    if employee_code:
        emp_qs = emp_qs.filter(employee_code=employee_code)
    elif code:
        emp_qs = emp_qs.filter(employee_code__icontains=code)
    if search:
        emp_qs = emp_qs.filter(
            Q(employee_code__icontains=search) | 
            Q(name__icontains=search)
        )

    pending_list = []
    for emp in emp_qs.order_by('employee_code'):
        pending_list.append({
            'id': f"pending_{emp.employee_code}_{emp.id}",
            'employee_code': emp.employee_code,
            'employee_name': emp.name,
            'department': emp.department,
            'location': emp.location or '',
            'branch': emp.branch.name if emp.branch else '',
            'designation': emp.designation,
            'date_of_joining': emp.date_of_joining.strftime('%Y-%m-%d') if emp.date_of_joining else '',
            'assignment_period': emp.assignment_period or '',
            'job_competence': 0,
            'productivity_responsibility': 0,
            'communication_teamwork': 0,
            'professionalism_discipline': 0,
            'initiative_improvement': 0,
            'performance_score': 0,
            'total_deduction': 0,
            'final_score': 0,
            'rating': 'Pending',
            'submitted_date': None,
            'submitted_by': None,
            'submitted_by_detail': None,
            'status': 'Pending',
            'is_pending': True
        })
    return pending_list


class AppraisalViewSet(viewsets.ModelViewSet):
    queryset = Appraisal.objects.all().order_by('-submitted_date')
    serializer_class = AppraisalSerializer

    def create(self, request, *args, **kwargs):
        user = request.user
        emp_code = str(request.data.get('employee_code') or '').strip()
        if user and user.is_authenticated and hasattr(user, 'department_admin'):
            if emp_code.lower() == user.username.strip().lower():
                return Response({"non_field_errors": ["Self-appraisal is not permitted. Department administrators cannot submit appraisals for themselves."]}, status=status.HTTP_400_BAD_REQUEST)
        return super().create(request, *args, **kwargs)

    def get_permissions(self):
        if self.action in ['create', 'update', 'partial_update', 'list', 'retrieve', 'approve']:
            return [permissions.AllowAny()]
        return [IsAdminUserOrReadOnly()]

    def list(self, request, *args, **kwargs):
        status_param = request.query_params.get('status', '').lower()
        if status_param == 'pending':
            records = get_pending_appraisal_records(request.user, request.query_params)
            return Response(records)
        elif status_param == 'all':
            queryset = self.filter_queryset(self.get_queryset())
            serializer = self.get_serializer(queryset, many=True)
            submitted_data = serializer.data
            pending_records = get_pending_appraisal_records(request.user, request.query_params)
            return Response(submitted_data + pending_records)
        else:
            return super().list(request, *args, **kwargs)

    def get_queryset(self):
        queryset = super().get_queryset()
        user = self.request.user
        is_admin = user and user.is_authenticated and (user.is_staff or user.is_superuser or hasattr(user, 'department_admin'))
        
        code_param = self.request.query_params.get('employee_code')

        if not is_admin:
            if code_param:
                queryset = queryset.filter(employee_code=code_param)
            elif self.kwargs.get('pk'):
                if user and user.is_authenticated:
                    queryset = queryset.filter(employee_code=user.username)
            else:
                queryset = queryset.none()
        else:
            # Check if user is a department_admin
            if hasattr(user, 'department_admin'):
                dept_admin = user.department_admin
                if not dept_admin.departments or not dept_admin.locations or not dept_admin.branches:
                    return queryset.none()
                queryset = queryset.exclude(employee_code=user.username)
                dept_filter = Q()
                for d in user.department_admin.departments:
                    dept_filter |= Q(department__iexact=d) | Q(department__icontains=d)
                queryset = queryset.filter(dept_filter)
                loc_filter = Q()
                for loc in user.department_admin.locations:
                    loc_filter |= build_exact_location_filter(loc)
                queryset = queryset.filter(loc_filter)
                br_filter = Q()
                for br in user.department_admin.branches:
                    br_filter |= Q(branch__name__iexact=br)
                queryset = queryset.filter(br_filter)

                # Restrict department admin to appraisals in the active period (unless querying specific employee history / all_periods)
                code_query = self.request.query_params.get('employee_code') or self.request.query_params.get('code')
                all_periods_flag = str(self.request.query_params.get('all_periods') or '').lower() in ['true', '1', 'yes']
                if not (code_query or all_periods_flag):
                    active_p = IncentivePeriod.objects.filter(is_active=True).first()
                    if active_p:
                        emp_codes_in_period = Employee.objects.filter(
                            Q(incentive_period=active_p) | 
                            Q(assignment_period__iexact=active_p.name) | 
                            Q(assignment_period__iexact=active_p.months)
                        ).values_list('employee_code', flat=True)

                        queryset = queryset.filter(
                            Q(assignment_period__iexact=active_p.name) | 
                            Q(assignment_period__iexact=active_p.months) |
                            Q(employee_code__in=emp_codes_in_period)
                        )

            # Query Filters
            dept = self.request.query_params.get('department')
            location_param = self.request.query_params.get('location')
            branch_param = self.request.query_params.get('branch')
            desig = self.request.query_params.get('designation')
            period = self.request.query_params.get('assignment_period')
            rating = self.request.query_params.get('rating')
            name = self.request.query_params.get('name')
            code = self.request.query_params.get('code')
            employee_code = self.request.query_params.get('employee_code')
            start_date = self.request.query_params.get('start_date')
            end_date = self.request.query_params.get('end_date')
            search = self.request.query_params.get('search')

            if dept:
                queryset = queryset.filter(Q(department__iexact=dept) | Q(department__icontains=dept))
            if location_param:
                queryset = queryset.filter(build_exact_location_filter(location_param))
            if branch_param:
                queryset = queryset.filter(branch__name__iexact=branch_param)
            if desig:
                queryset = queryset.filter(designation__iexact=desig)
            if period:
                queryset = queryset.filter(assignment_period=period)
            if rating:
                queryset = queryset.filter(rating__icontains=rating)
            if name:
                queryset = queryset.filter(employee_name__icontains=name)
            if employee_code:
                clean_code = str(employee_code).strip()
                code_q = Q(employee_code__iexact=clean_code)
                if clean_code.isdigit():
                    code_q |= Q(employee_code__iexact=clean_code.zfill(6))
                queryset = queryset.filter(code_q)
            elif code:
                clean_code = str(code).strip()
                code_q = Q(employee_code__iexact=clean_code) | Q(employee_code__icontains=clean_code)
                if clean_code.isdigit():
                    code_q |= Q(employee_code__iexact=clean_code.zfill(6))
                queryset = queryset.filter(code_q)
            if start_date:
                queryset = queryset.filter(submitted_date__date__gte=start_date)
            if end_date:
                queryset = queryset.filter(submitted_date__date__lte=end_date)
            if search:
                clean_s = str(search).strip()
                if clean_s:
                    s_q = Q(employee_code__icontains=clean_s) | Q(employee_name__icontains=clean_s)
                    if clean_s.isdigit():
                        s_q |= Q(employee_code__iexact=clean_s.zfill(6))
                    queryset = queryset.filter(s_q)

        return queryset

    def perform_create(self, serializer):
        user = self.request.user if self.request.user.is_authenticated else None
        serializer.save(submitted_by=user)

    def perform_update(self, serializer):
        instance = self.get_object()
        if instance.status == 'Approved':
            raise permissions.exceptions.PermissionDenied("Approved appraisals cannot be modified.")
        serializer.save()

    @action(detail=True, methods=['post'])
    def approve(self, request, pk=None):
        appraisal = self.get_object()
        if appraisal.status == 'Approved':
            return Response({"detail": "Appraisal is already approved."}, status=status.HTTP_400_BAD_REQUEST)
        appraisal.status = 'Approved'
        appraisal.save()
        return Response({"status": "Appraisal approved successfully."})


class ExcelExportView(APIView):
    permission_classes = [IsAdminUserOrReadOnly]

    def get(self, request):
        status_param = request.query_params.get('status', '').lower()

        rows_to_export = []

        if status_param != 'pending':
            view = AppraisalViewSet()
            view.request = request
            queryset = view.get_queryset()
            for obj in queryset:
                rows_to_export.append({
                    'code': obj.employee_code,
                    'name': obj.employee_name,
                    'dept': obj.department,
                    'loc': obj.location,
                    'branch': obj.branch.name if obj.branch else '',
                    'desig': obj.designation,
                    'doj': obj.date_of_joining.strftime('%Y-%m-%d') if obj.date_of_joining else '',
                    'period': obj.assignment_period,
                    'perf_score': obj.performance_score,
                    'deduction': obj.total_deduction,
                    'final_score': obj.final_score,
                    'rating': obj.rating,
                    'date': obj.submitted_date.strftime('%Y-%m-%d %H:%M') if obj.submitted_date else '',
                    'by': obj.submitted_by.username if obj.submitted_by else '-'
                })

        if status_param in ['pending', 'all']:
            pending_records = get_pending_appraisal_records(request.user, request.query_params)
            for rec in pending_records:
                rows_to_export.append({
                    'code': rec['employee_code'],
                    'name': rec['employee_name'],
                    'dept': rec['department'],
                    'loc': rec['location'],
                    'branch': rec.get('branch', ''),
                    'desig': rec['designation'],
                    'doj': rec['date_of_joining'],
                    'period': rec['assignment_period'],
                    'perf_score': '-',
                    'deduction': '-',
                    'final_score': '-',
                    'rating': 'Pending',
                    'date': '-',
                    'by': '-'
                })

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Appraisal Reports"

        headers = [
            "Employee Code", "Employee Name", "Department", "Location", "Branch", "Designation", 
            "Date of Joining", "Assignment Period", "Performance Score", 
            "Deduction Score", "Final Score", "Rating", "Submitted Date", "Done By (Admin)"
        ]
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

        for row_num, item in enumerate(rows_to_export, 2):
            ws.cell(row=row_num, column=1, value=item['code'])
            ws.cell(row=row_num, column=2, value=item['name'])
            ws.cell(row=row_num, column=3, value=item['dept'])
            ws.cell(row=row_num, column=4, value=item['loc'])
            ws.cell(row=row_num, column=5, value=item['branch'])
            ws.cell(row=row_num, column=6, value=item['desig'])
            ws.cell(row=row_num, column=7, value=item['doj'])
            ws.cell(row=row_num, column=8, value=item['period'])
            ws.cell(row=row_num, column=9, value=item['perf_score'])
            ws.cell(row=row_num, column=10, value=item['deduction'])
            ws.cell(row=row_num, column=11, value=item['final_score'])
            ws.cell(row=row_num, column=12, value=item['rating'])
            ws.cell(row=row_num, column=13, value=item['date'])
            ws.cell(row=row_num, column=14, value=item['by'])

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 12)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = "attachment; filename=appraisal_reports.xlsx"
        wb.save(response)
        return response


class DashboardStatsView(APIView):
    permission_classes = [IsAdminUserOrReadOnly]

    def get(self, request):
        user = request.user
        branch_param = request.query_params.get('branch')
        
        # Scope strictly by active incentive period
        active_period = IncentivePeriod.objects.filter(is_active=True).first()

        employees = Employee.objects.filter(is_trashed=False)
        appraisals = Appraisal.objects.all()

        if active_period:
            employees = employees.filter(
                Q(incentive_period=active_period) |
                Q(assignment_period__iexact=active_period.name) |
                Q(assignment_period__iexact=active_period.months)
            )
            appraisals = appraisals.filter(
                Q(assignment_period__iexact=active_period.name) |
                Q(assignment_period__iexact=active_period.months)
            )

        # Exclude trashed staff from appraisals
        trashed_codes = list(Employee.objects.filter(is_trashed=True).values_list('employee_code', flat=True))
        if trashed_codes:
            appraisals = appraisals.exclude(employee_code__in=trashed_codes)

        # Department admin access restrictions
        dept_admin_info = None
        if hasattr(user, 'department_admin'):
            dept_admin = user.department_admin

            if not dept_admin.departments or not dept_admin.locations or not dept_admin.branches:
                employees = employees.none()
                appraisals = appraisals.none()
            else:
                dept_filter = Q()
                for d in dept_admin.departments:
                    dept_filter |= Q(department__iexact=d) | Q(department__icontains=d)
                employees = employees.filter(dept_filter).exclude(employee_code=user.username)
                appraisals = appraisals.filter(dept_filter).exclude(employee_code=user.username)
                loc_filter = Q()
                for loc in dept_admin.locations:
                    loc_filter |= build_exact_location_filter(loc)
                employees = employees.filter(loc_filter)
                appraisals = appraisals.filter(loc_filter)
                br_filter = Q()
                for br in dept_admin.branches:
                    br_filter |= Q(branch__name__iexact=br)
                employees = employees.filter(br_filter)
                appraisals = appraisals.filter(br_filter)

            dept_admin_info = {
                'username': user.username,
                'departments': dept_admin.departments,
                'locations': dept_admin.locations,
                'branches': dept_admin.branches,
                'email': dept_admin.effective_email,
                'whatsapp_number': dept_admin.whatsapp_number,
            }

        if branch_param:
            employees = employees.filter(branch__name__iexact=branch_param)
            appraisals = appraisals.filter(branch__name__iexact=branch_param)

        total_employees = employees.count()
        total_appraisals = appraisals.count()
        completion_rate = round((total_appraisals / total_employees * 100), 1) if total_employees > 0 else 0
        
        # Pending appraisals = Active Employees in this active subset who have not submitted any appraisal
        submitted_codes = appraisals.values_list('employee_code', flat=True).distinct()
        pending_employees_qs = employees.filter(status='Active').exclude(employee_code__in=submitted_codes).order_by('employee_code')
        pending_appraisals = pending_employees_qs.count()
        pending_list = [
            {
                'employee_code': emp.employee_code,
                'name': emp.name,
                'department': emp.department,
                'designation': emp.designation,
                'location': emp.location,
                'branch': emp.branch.name if emp.branch else '',
                'date_of_joining': emp.date_of_joining.strftime('%d %b %Y') if emp.date_of_joining else '',
                'assignment_period': emp.assignment_period or '',
                'category': emp.category or '',
                'level': emp.level or '',
                'division': emp.division or '',
                'status': emp.status
            }
            for emp in pending_employees_qs.select_related('branch')
        ]

        submitted_list = [
            {
                'id': app.id,
                'employee_code': app.employee_code,
                'name': app.employee_name,
                'department': app.department,
                'designation': app.designation,
                'location': app.location,
                'branch': app.branch.name if app.branch else '',
                'date_of_joining': app.date_of_joining.strftime('%d %b %Y') if app.date_of_joining else '',
                'assignment_period': app.assignment_period or '',
                'job_competence': app.job_competence,
                'productivity_responsibility': app.productivity_responsibility,
                'communication_teamwork': app.communication_teamwork,
                'professionalism_discipline': app.professionalism_discipline,
                'initiative_improvement': app.initiative_improvement,
                'performance_score': app.performance_score,
                'total_deduction': app.total_deduction,
                'marks_obtained': app.marks_obtained,
                'final_score': round(float(app.final_score or 0.0), 2),
                'rating': app.rating,
                'submitted_date': app.submitted_date.strftime('%d %b %Y, %I:%M %p') if app.submitted_date else '',
                'status': app.status,
            }
            for app in appraisals.select_related('branch').order_by('-submitted_date', '-id')
        ]

        # Score tiers on 3.0 scale (replacing obsolete letter grade system)
        tier1_count = appraisals.filter(final_score__gte=2.70).count()
        tier2_count = appraisals.filter(final_score__gte=2.25, final_score__lt=2.70).count()
        tier3_count = appraisals.filter(final_score__gte=1.80, final_score__lt=2.25).count()
        tier4_count = appraisals.filter(final_score__gte=1.50, final_score__lt=1.80).count()
        tier5_count = appraisals.filter(final_score__lt=1.50).count()

        score_distribution = [
            {
                'key': 'tier1',
                'label': 'Score 2.70 – 3.00 (90% – 100%)',
                'short_label': '2.70 – 3.00',
                'percentage_range': '90% – 100%',
                'count': tier1_count,
                'color': '#00b87c'
            },
            {
                'key': 'tier2',
                'label': 'Score 2.25 – 2.69 (75% – 89%)',
                'short_label': '2.25 – 2.69',
                'percentage_range': '75% – 89%',
                'count': tier2_count,
                'color': '#0066ff'
            },
            {
                'key': 'tier3',
                'label': 'Score 1.80 – 2.24 (60% – 74%)',
                'short_label': '1.80 – 2.24',
                'percentage_range': '60% – 74%',
                'count': tier3_count,
                'color': '#8b5cf6'
            },
            {
                'key': 'tier4',
                'label': 'Score 1.50 – 1.79 (50% – 59%)',
                'short_label': '1.50 – 1.79',
                'percentage_range': '50% – 59%',
                'count': tier4_count,
                'color': '#f97316'
            },
            {
                'key': 'tier5',
                'label': 'Score < 1.50 (Below 50%)',
                'short_label': '< 1.50',
                'percentage_range': 'Below 50%',
                'count': tier5_count,
                'color': '#ef4444'
            },
        ]

        rating_counts = {
            'tier1': tier1_count,
            'tier2': tier2_count,
            'tier3': tier3_count,
            'tier4': tier4_count,
            'tier5': tier5_count,
            'A': tier1_count,
            'B': tier2_count,
            'C': tier3_count,
            'D': tier4_count,
            'E': tier5_count,
        }

        # Department-wise appraisal count and average final score (strictly from current period appraisals)
        dept_stats_qs = appraisals.values('department').annotate(
            submitted=Count('id'),
            avg_score=Avg('final_score')
        ).order_by('-submitted', 'department')
        
        department_stats = []
        for ds in dept_stats_qs:
            department_stats.append({
                'department': ds['department'],
                'submitted': ds['submitted'],
                'avg_score': round(float(ds['avg_score'] or 0.0), 2)
            })

        return Response({
            'total_employees': total_employees,
            'total_appraisals': total_appraisals,
            'pending_appraisals': pending_appraisals,
            'completion_rate': completion_rate,
            'ratings': rating_counts,
            'score_distribution': score_distribution,
            'department_stats': department_stats,
            'pending_list': pending_list,
            'submitted_list': submitted_list,
            'dept_admin_info': dept_admin_info,
            'active_period': active_period.name if active_period else None
        }, status=status.HTTP_200_OK)


class TrainingEmployeeViewSet(viewsets.ModelViewSet):
    queryset = TrainingEmployee.objects.all().order_by('-date_of_joining', 'employee_code')
    serializer_class = TrainingEmployeeSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        queryset = super().get_queryset()
        user = self.request.user
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            queryset = queryset.filter(department__in=user.department_admin.departments)
            if user.department_admin.locations:
                loc_filter = Q()
                for loc in user.department_admin.locations:
                    loc_filter |= build_exact_location_filter(loc)
                queryset = queryset.filter(loc_filter)
            if user.department_admin.branches:
                br_filter = Q()
                for br in user.department_admin.branches:
                    br_filter |= Q(branch__name__iexact=br)
                queryset = queryset.filter(br_filter)

        dept = self.request.query_params.get('department')
        branch_param = self.request.query_params.get('branch')
        status_param = self.request.query_params.get('status')
        search = self.request.query_params.get('search')
        due_only = self.request.query_params.get('due_only')

        if dept:
            queryset = queryset.filter(department__iexact=dept)
        if branch_param:
            queryset = queryset.filter(branch__name__iexact=branch_param)
        if status_param and status_param != 'All':
            queryset = queryset.filter(status=status_param)
        if search:
            queryset = queryset.filter(
                Q(employee_code__icontains=search) |
                Q(name__icontains=search)
            )
        if due_only == 'true':
            import datetime
            today = datetime.date.today()
            target_date = today + datetime.timedelta(days=7)
            queryset = queryset.filter(
                training_end_date__lte=target_date,
                whatsapp_notification_sent=False,
                status__in=['Training', 'Assessment Due']
            )

        return queryset

    def create(self, request, *args, **kwargs):
        emp_code = str(request.data.get('employee_code') or '').strip()
        name = str(request.data.get('name') or '').strip()
        dept = str(request.data.get('department') or '').strip()

        if not emp_code:
            return Response({"employee_code": ["Employee Code is required."]}, status=status.HTTP_400_BAD_REQUEST)
        if not name:
            return Response({"name": ["Employee Name is required."]}, status=status.HTTP_400_BAD_REQUEST)
        if not dept:
            return Response({"department": ["Department is required."]}, status=status.HTTP_400_BAD_REQUEST)

        # Department admin permission check
        user = request.user
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            allowed_depts = [d.strip().lower() for d in user.department_admin.departments]
            if dept.lower() not in allowed_depts:
                return Response({"department": [f"You do not have permission to add trainees in department '{dept}'."]}, status=status.HTTP_403_FORBIDDEN)

        existing = TrainingEmployee.objects.filter(employee_code=emp_code).first()
        if existing:
            serializer = self.get_serializer(existing, data=request.data, partial=True)
        else:
            serializer = self.get_serializer(data=request.data)

        serializer.is_valid(raise_exception=True)
        trainee = serializer.save()

        headers = self.get_success_headers(serializer.data)
        return Response(serializer.data, status=status.HTTP_201_CREATED if not existing else status.HTTP_200_OK, headers=headers)


    @action(detail=True, methods=['post'], url_path='send-email')
    def send_email(self, request, pk=None):
        trainee = self.get_object()
        custom_base_url = request.data.get('base_url')
        notify_candidates = request.data.get('notify_candidates', False)
        res = send_department_wise_trainee_emails(
            [trainee],
            is_reminder=False,
            custom_base_url=custom_base_url,
            notify_candidates=notify_candidates
        )
        trainee.refresh_from_db()
        serializer = self.get_serializer(trainee)
        success = res["sent_staff"] > 0
        message = res["message"]
        dept_result = res["department_results"][0] if res["department_results"] else {}
        if success:
            return Response({
                "success": True,
                "message": message,
                "trainee": serializer.data,
                "total_staff": 1,
                "sent_staff": 1,
                "failed_staff": 0,
                "total_departments": 1,
                "sent_departments": 1,
                "failed_departments": 0,
                "department_results": res["department_results"],
                "details": res["details"]
            })
        else:
            return Response({
                "success": False,
                "error": message,
                "trainee": serializer.data,
                "total_staff": 1,
                "sent_staff": 0,
                "failed_staff": 1,
                "total_departments": 1,
                "sent_departments": 0,
                "failed_departments": 1,
                "department_results": res["department_results"],
                "details": res["details"]
            }, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='send-whatsapp')
    def send_whatsapp(self, request, pk=None):
        # WhatsApp notifications temporarily redirected to Email
        return self.send_email(request, pk=pk)

    @action(detail=False, methods=['post'], url_path='send-bulk-email')
    def send_bulk_email(self, request):
        employee_ids = request.data.get('employee_ids', [])
        if isinstance(employee_ids, str):
            import json
            try:
                employee_ids = json.loads(employee_ids)
            except Exception:
                employee_ids = [x.strip() for x in employee_ids.split(',') if x.strip()]
        elif not isinstance(employee_ids, list):
            employee_ids = request.data.getlist('employee_ids') if hasattr(request.data, 'getlist') else [employee_ids]

        employee_ids = [x for x in employee_ids if x]
        if not employee_ids:
            return Response({"error": "List of employee_ids is required."}, status=status.HTTP_400_BAD_REQUEST)

        trainees = list(TrainingEmployee.objects.filter(id__in=employee_ids))
        if not trainees:
            return Response({"error": "No trainees found for the provided IDs."}, status=status.HTTP_404_NOT_FOUND)

        custom_base_url = request.data.get('base_url')
        notify_candidates = request.data.get('notify_candidates', False)

        result = send_department_wise_trainee_emails(
            trainees,
            is_reminder=False,
            custom_base_url=custom_base_url,
            notify_candidates=notify_candidates
        )

        return Response(result, status=status.HTTP_200_OK if result["failed_staff"] == 0 else status.HTTP_207_MULTI_STATUS)

    @action(detail=False, methods=['post'], url_path='send-pending-by-department')
    def send_pending_by_department(self, request):
        """
        Scans all pending trainees (status in Training or Assessment Due) and dispatches
        ONE email per department to each Department Head with the table of pending staff.
        """
        department = request.data.get('department')
        custom_base_url = request.data.get('base_url')
        notify_candidates = request.data.get('notify_candidates', False)
        only_unsent = request.data.get('only_unsent', True)

        qs = TrainingEmployee.objects.filter(status__in=['Training', 'Assessment Due'])
        if department and department.strip() and department.strip() != 'All':
            qs = qs.filter(department__iexact=department.strip())

        if only_unsent:
            qs = qs.filter(email_notification_sent=False)

        trainees = list(qs)
        if not trainees:
            return Response({
                "success": True,
                "message": "No pending staff found requiring email intimation for the selected criteria.",
                "total_staff": 0,
                "sent_staff": 0,
                "failed_staff": 0,
                "total_departments": 0,
                "sent_departments": 0,
                "failed_departments": 0,
                "department_results": [],
                "total": 0,
                "sent": 0,
                "failed": 0,
                "details": []
            }, status=status.HTTP_200_OK)

        result = send_department_wise_trainee_emails(
            trainees,
            is_reminder=False,
            custom_base_url=custom_base_url,
            notify_candidates=notify_candidates
        )
        return Response(result, status=status.HTTP_200_OK if result["failed_staff"] == 0 else status.HTTP_207_MULTI_STATUS)

    @action(detail=False, methods=['post'], url_path='send-all-reminders')
    def send_all_reminders(self, request):
        custom_base_url = request.data.get('base_url')
        result = check_and_send_all_due_email_reminders(custom_base_url=custom_base_url)
        return Response(result, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='sample-template')
    def sample_template(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Trainee Upload Template"

        headers = [
            "Employee Code", "Employee Name", "Department", "Location", "Branch", 
            "Designation", "Date of Joining (YYYY-MM-DD)", "Candidate Email", "Training Period"
        ]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

        sample_rows = [
            ["TRN001", "Amina Patel", "Nursing", "Floor 2", "IQRAA INTERNATIONAL HOSPITAL & RESEARCH CENTRE", "Trainee Staff Nurse", "2026-08-25", "amina.patel@iqraahospital.in", "1 Month"],
            ["TRN002", "Rahul Sharma", "Pharmacy", "Main Pharmacy", "IQRAA INTERNATIONAL HOSPITAL & RESEARCH CENTRE", "Junior Pharmacist Trainee", "2026-08-30", "rahul.sharma@iqraahospital.in", "6 Months"],
        ]
        for row_num, row_vals in enumerate(sample_rows, 2):
            for col_num, val in enumerate(row_vals, 1):
                ws.cell(row=row_num, column=col_num, value=val)

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 16)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=trainee_upload_template.xlsx"
        wb.save(response)
        return response

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        queryset = self.filter_queryset(self.get_queryset())
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Trainee Master"

        headers = [
            "Employee Code", "Employee Name", "Department", "Location", "Branch", 
            "Designation", "Date of Joining", "Training Period", "Training End Date", 
            "Candidate Email", "Department Email", "Status", 
            "Email Sent", "Email Date", "Email Status"
        ]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

        for row_num, obj in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=obj.employee_code)
            ws.cell(row=row_num, column=2, value=obj.name)
            ws.cell(row=row_num, column=3, value=obj.department)
            ws.cell(row=row_num, column=4, value=obj.location)
            ws.cell(row=row_num, column=5, value=obj.branch.name if obj.branch else '')
            ws.cell(row=row_num, column=6, value=obj.designation)
            ws.cell(row=row_num, column=7, value=obj.date_of_joining.strftime('%Y-%m-%d') if obj.date_of_joining else '')
            ws.cell(row=row_num, column=8, value=f"{obj.training_period_months} Month(s)")
            ws.cell(row=row_num, column=9, value=obj.training_end_date.strftime('%Y-%m-%d') if obj.training_end_date else '')
            ws.cell(row=row_num, column=10, value=obj.email)
            ws.cell(row=row_num, column=11, value=obj.department_email)
            ws.cell(row=row_num, column=12, value=obj.status)
            ws.cell(row=row_num, column=13, value="Yes" if obj.email_notification_sent else "No")
            ws.cell(row=row_num, column=14, value=obj.email_notification_date.strftime('%Y-%m-%d %H:%M') if obj.email_notification_date else '')
            ws.cell(row=row_num, column=15, value=obj.email_notification_status)

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 14)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=trainees_list.xlsx"
        wb.save(response)
        return response


class TrainingExcelImportView(APIView):
    permission_classes = [IsAdminUserOrReadOnly]

    def post(self, request):
        if 'file' not in request.FILES:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        excel_file = request.FILES['file']
        try:
            headers, rows = extract_tabular_data(excel_file)
        except Exception as e:
            return Response({"error": f"Invalid excel file format: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)

        col_mapping = {
            'employee_code': ['employee code', 'code', 'emp code', 'emp_code', 'employee_code', 'emp id', 'empid', 'emp_id', 'employee id', 'staff id', 'staff code', 'staff no', 'trainee code', 'token no', 'id'],
            'name': ['employee name', 'name', 'emp name', 'emp_name', 'employee_name', 'staff name', 'trainee name', 'full name', 'staff'],
            'department': ['department', 'dept', 'department name', 'dept name', 'dept.', 'department_name'],
            'location': ['location', 'loc', 'location name', 'ward', 'area', 'place', 'work location', 'duty station', 'station'],
            'branch': ['branch', 'brach', 'hospital branch', 'unit', 'hospital unit', 'hospital'],
            'designation': ['designation', 'desig', 'role', 'post', 'position', 'job title', 'title'],
            'date_of_joining': ['date of joining', 'joining date', 'doj', 'joining_date', 'date_of_joining', 'join date', 'date of join', 'd.o.j', 'training start date', 'start date'],
            'email': ['candidate email', 'candidate_email', 'trainee email', 'email id', 'email_id', 'email', 'candidate mail', 'mail', 'staff email'],
            'training_period': ['training period', 'training_period', 'period', 'duration', 'training duration', 'months', 'training period months', 'period months']
        }

        indices = get_column_indices(headers, col_mapping)
        required_keys = ['employee_code', 'name', 'department']
        for rk in required_keys:
            if rk not in indices:
                return Response({
                    "error": f"Missing required column matching: {rk.replace('_', ' ').title()}. Please check Excel headers."
                }, status=status.HTTP_400_BAD_REQUEST)

        selected_branch = request.data.get('selected_branch') or request.query_params.get('branch') or 'IQRAA INTERNATIONAL HOSPITAL & RESEARCH CENTRE'

        # -------------------------------------------------------------
        # VALIDATION: Check if Department Admin email is added in Admin Login
        # If any department in the uploaded file does not have a configured admin email,
        # reject the ENTIRE file without importing any rows.
        # -------------------------------------------------------------
        departments_in_file = set()
        for row_data in rows:
            if not any(row_data):
                continue
            if 'department' in indices and indices['department'] < len(row_data) and row_data[indices['department']]:
                d_val = str(row_data[indices['department']]).strip()
                if d_val:
                    departments_in_file.add(d_val)

        missing_dept_emails = []
        for dept_name in sorted(departments_in_file):
            dept_emails = get_department_emails(dept_name)
            if not dept_emails:
                missing_dept_emails.append(dept_name)

        if missing_dept_emails:
            missing_str = ", ".join(missing_dept_emails)
            return Response({
                "error": f"Upload rejected: Department email ID is not added in Admin Login for department(s): {missing_str}. Please add the department email ID in Admin Login before uploading the file."
            }, status=status.HTTP_400_BAD_REQUEST)

        success_count = 0
        skipped_count = 0
        failed_rows = []
        row_num = 1

        for row_data in rows:
            row_num += 1
            if not any(row_data):
                continue

            try:
                raw_code = row_data[indices['employee_code']] if indices.get('employee_code') is not None and indices['employee_code'] < len(row_data) else ""
                if isinstance(raw_code, float) and raw_code.is_integer():
                    emp_code = str(int(raw_code)).strip()
                else:
                    emp_code = str(raw_code or "").strip()
                    if emp_code.endswith('.0') and emp_code[:-2].replace('-', '').isdigit():
                        emp_code = emp_code[:-2]
                name = str(row_data[indices['name']]).strip() if indices.get('name') is not None and row_data[indices['name']] else ""
                dept = str(row_data[indices['department']]).strip() if indices.get('department') is not None and row_data[indices['department']] else ""
                location = str(row_data[indices['location']]).strip() if 'location' in indices and indices['location'] < len(row_data) and row_data[indices['location']] else ""
                branch_str = str(row_data[indices['branch']]).strip() if 'branch' in indices and indices['branch'] < len(row_data) and row_data[indices['branch']] else ""
                desig = str(row_data[indices['designation']]).strip() if 'designation' in indices and indices['designation'] < len(row_data) and row_data[indices['designation']] else "Trainee Staff"
                doj_raw = row_data[indices['date_of_joining']] if 'date_of_joining' in indices and indices['date_of_joining'] < len(row_data) else None
                cand_email = str(row_data[indices['email']]).strip() if 'email' in indices and indices['email'] < len(row_data) and row_data[indices['email']] else ""
                
                # Parse period (1 Month or 6 Months; default 1)
                period_raw = str(row_data[indices['training_period']]).strip() if 'training_period' in indices and indices['training_period'] < len(row_data) and row_data[indices['training_period']] else ""
                period_months = 1
                if '6' in period_raw:
                    period_months = 6
                elif '1' in period_raw:
                    period_months = 1
            except IndexError:
                failed_rows.append({"row": row_num, "code": "", "reason": "Missing cell values matching header"})
                continue

            if not emp_code or not name:
                failed_rows.append({"row": row_num, "code": emp_code, "reason": "Employee Code and Name cannot be empty"})
                continue

            # Parse Date of Joining
            doj = parse_excel_date(doj_raw)

            if not branch_str:
                branch_str = selected_branch

            branch_obj = None
            if branch_str:
                branch_obj, _ = Branch.objects.get_or_create(name=branch_str)

            # Auto-resolve department email from registered Department Admin
            dept_emails = get_department_emails(dept)
            dept_email_str = dept_emails[0] if dept_emails else ""

            try:
                # Create or update TrainingEmployee record
                trainee, created = TrainingEmployee.objects.update_or_create(
                    employee_code=emp_code,
                    defaults={
                        'name': name,
                        'department': dept,
                        'location': location,
                        'branch': branch_obj,
                        'designation': desig or 'Trainee Staff',
                        'email': cand_email,
                        'department_email': dept_email_str,
                        'date_of_joining': doj,
                        'training_period_months': period_months,
                        'training_end_date': None,  # Will auto-calculate on save
                        'status': 'Training'
                    }
                )

                # Sync to Employee table with status='Training' so Department Admins see them
                # and EmployeeAutoFetchView can resolve them for appraisal forms
                emp_record, _ = Employee.objects.update_or_create(
                    employee_code=emp_code,
                    department=dept,
                    location=location,
                    defaults={
                        'name': name,
                        'branch': branch_obj,
                        'designation': desig or 'Trainee Staff',
                        'date_of_joining': doj,
                        'assignment_period': f"{period_months}-Month Training Period",
                        'status': 'Training'
                    }
                )

                # Auto-create User account if not exists with default password Iqraa@123
                if not User.objects.filter(username=emp_code).exists():
                    User.objects.create_user(username=emp_code, password='Iqraa@123', is_staff=False)

                success_count += 1
            except Exception as ex:
                failed_rows.append({"row": row_num, "code": emp_code, "reason": str(ex)})

        summary = {
            "total_processed": len(rows),
            "successfully_imported": success_count,
            "skipped_duplicates": skipped_count,
            "failed": len(failed_rows)
        }
        return Response({
            "summary": summary,
            "failed_rows": failed_rows
        }, status=status.HTTP_200_OK if len(failed_rows) == 0 else status.HTTP_207_MULTI_STATUS)


class WhatsAppConfigView(APIView):
    permission_classes = [permissions.IsAdminUser]

    def get(self, request):
        config = get_whatsapp_config()
        serializer = WhatsAppConfigSerializer(config)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def post(self, request):
        action_type = request.data.get('action')
        if action_type == 'test':
            to_phone = request.data.get('phone_number')
            if not to_phone:
                return Response({"error": "Phone number is required for test message."}, status=status.HTTP_400_BAD_REQUEST)
            config = get_whatsapp_config()
            test_text = f"Hello! This is a test notification from Employee Appraisal & Training System. Meta WhatsApp Cloud API is successfully connected! (Time: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')})"
            success, result = send_meta_whatsapp_message(
                to_number=to_phone,
                force_text=True,
                fallback_text=test_text
            )
            if success:
                return Response({"success": True, "message": f"Test message delivered successfully to {to_phone}!"})
            else:
                return Response({"success": False, "error": f"Delivery failed: {result}"}, status=status.HTTP_400_BAD_REQUEST)

        # Update settings
        config = get_whatsapp_config()
        serializer = WhatsAppConfigSerializer(config, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_200_OK)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class EmailConfigView(APIView):
    permission_classes = [permissions.IsAdminUser]

    def get(self, request):
        config = get_email_config()
        serializer = EmailConfigSerializer(config)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def post(self, request):
        action_type = request.data.get('action')
        config = get_email_config()

        if action_type == 'test':
            to_email = request.data.get('to_email') or request.data.get('email')
            if not to_email:
                return Response({"error": "Recipient email address is required for test message."}, status=status.HTTP_400_BAD_REQUEST)

            test_subject = "Test Notification: Employee Performance Appraisal & Training System"
            test_text = f"Hello!\n\nThis is a test email sent from the Employee Performance Appraisal & Training System (IQRAA International Hospital & Research Centre).\n\nSMTP Host: {config.smtp_host}:{config.smtp_port}\nSender: {config.from_email}\nTimestamp: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n\nYour SMTP email delivery is operating properly!"
            test_html = f"""<div style="font-family: Arial, sans-serif; padding: 20px; color: #1e293b;">
              <h2 style="color: #0f766e;">Email Configuration Test</h2>
              <p>This is a test notification from the <strong>Employee Performance Appraisal & Training System</strong>.</p>
              <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 12px; margin: 14px 0; font-size: 13px;">
                <strong>Sender:</strong> {config.from_email}<br>
                <strong>SMTP Server:</strong> {config.smtp_host}:{config.smtp_port}<br>
                <strong>Security:</strong> {'TLS (STARTTLS)' if config.use_tls else ('SSL' if config.use_ssl else 'Plain')}<br>
                <strong>Time:</strong> {datetime.datetime.now().strftime('%d/%m/%Y %H:%M:%S')}
              </div>
              <p style="color: #059669; font-weight: bold;">✔ SMTP email service is working properly!</p>
              <hr style="border: 0; border-top: 1px solid #e2e8f0; margin: 18px 0;">
              <p style="font-size: 11px; color: #64748b;">IQRAA International Hospital & Research Centre</p>
            </div>"""

            success, msg = send_raw_email(test_subject, test_text, test_html, [to_email])
            if success:
                return Response({"success": True, "message": f"Test email delivered successfully to {to_email}!"})
            else:
                return Response({"success": False, "error": f"Email delivery failed: {msg}"}, status=status.HTTP_400_BAD_REQUEST)

        # Update email settings
        serializer = EmailConfigSerializer(config, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_200_OK)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class DepartmentLocationMetaView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        user = request.user

        # Fetch distinct non-empty departments
        emp_depts = Employee.objects.exclude(department='').values_list('department', flat=True).distinct()
        trn_depts = TrainingEmployee.objects.exclude(department='').values_list('department', flat=True).distinct()
        loc_depts = DutyLocation.objects.exclude(department='').values_list('department', flat=True).distinct()
        dept_set = set()
        for d in emp_depts:
            if d:
                for sub in d.split(','):
                    if sub.strip():
                        dept_set.add(sub.strip())
        for d in trn_depts:
            if d:
                for sub in d.split(','):
                    if sub.strip():
                        dept_set.add(sub.strip())
        for d in loc_depts:
            if d:
                for sub in d.split(','):
                    if sub.strip():
                        dept_set.add(sub.strip())
        all_departments = sorted(list(dept_set))

        # Fetch distinct non-empty locations (combining Master DutyLocation, Employee, and TrainingEmployee)
        emp_locs = Employee.objects.exclude(location='').values_list('location', flat=True).distinct()
        trn_locs = TrainingEmployee.objects.exclude(location='').values_list('location', flat=True).distinct()
        master_locs = DutyLocation.objects.filter(is_active=True).values_list('name', flat=True).distinct()
        loc_set = set()
        for l in emp_locs:
            if l:
                for sub in l.split(','):
                    if sub.strip():
                        loc_set.add(sub.strip())
        for l in trn_locs:
            if l:
                for sub in l.split(','):
                    if sub.strip():
                        loc_set.add(sub.strip())
        for l in master_locs:
            if l and l.strip():
                loc_set.add(l.strip())
        all_locations = sorted(list(loc_set))

        # Distinct organizations
        emp_orgs = Employee.objects.exclude(organization='').values_list('organization', flat=True).distinct()
        branch_orgs = Branch.objects.values_list('name', flat=True)
        default_orgs = [
            'IQRAA INTERNATIONAL HOSPITAL & RESEARCH CENTRE',
            'IQRAA HOSPITAL',
            'IQRAA COMMUNITY HOSPITAL -VAZHAKKAD',
            'RAHMA IQRAA COMMUNITY HOSPITAL-THOTTILPALAM',
            'MSS IQRAA COMMUNITY CLINIC- CHAKKUMKADAVU',
            'IQRAA HOSPITAL-SULTHAN BATHERY',
            'MSS IQRAA MEDICAL CENTRE-ADIMALI',
            'MEERA IQRAA SPECIALITY CLINIC-SULTHAN BATHERY',
            'IQRAA COMMUNITY CLINIC-KINASSERY',
            'IQRAA COMMUNITY CLINIC-PALAZHI',
            'IQRAA SPECIALITY CLINICS - ADDRESS MALL',
            'IQRAA COMMUNITY CLINIC-KODUVALLY',
            'IQRAA MATERNITY HOSPITAL - VAZHAKKAD',
            'IQRAA CENTRE FOR RESEARCH AND DEVELOPMENT',
            'IIM K - DISPENSARY'
        ]
        all_organizations = sorted(list(set(
            [o.strip() for o in emp_orgs if o and o.strip()] + 
            [b.strip() for b in branch_orgs if b and b.strip()] + 
            default_orgs
        )))

        # Build department -> locations mapping
        dept_locations = {}
        for emp in Employee.objects.exclude(department='').exclude(location='').values('department', 'location'):
            d = emp['department'].strip()
            l = emp['location'].strip()
            if d and l:
                d_subs = [x.strip() for x in d.split(',') if x.strip()]
                l_subs = [x.strip() for x in l.split(',') if x.strip()]
                for d_item in d_subs:
                    if d_item not in dept_locations:
                        dept_locations[d_item] = set()
                    for l_item in l_subs:
                        dept_locations[d_item].add(l_item)

        for mloc in DutyLocation.objects.filter(is_active=True).exclude(department='').values('department', 'name'):
            d = mloc['department'].strip()
            l = mloc['name'].strip()
            if d and l:
                if d not in dept_locations:
                    dept_locations[d] = set()
                dept_locations[d].add(l)

        dept_locations_map = {d: sorted(list(locs)) for d, locs in dept_locations.items()}

        allowed_departments = all_departments
        allowed_locations = all_locations

        # Department admin restrictions
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            da = user.department_admin
            if da.departments:
                allowed_departments = [d for d in all_departments if any(ad.lower() == d.lower() for ad in da.departments)]
                if da.locations:
                    allowed_locations = da.locations
                else:
                    loc_pool = set()
                    for d in allowed_departments:
                        loc_pool.update(dept_locations_map.get(d, []))
                    if loc_pool:
                        allowed_locations = sorted(list(loc_pool))

        return Response({
            'departments': allowed_departments,
            'all_departments': all_departments,
            'locations': allowed_locations,
            'all_locations': all_locations,
            'department_locations': dept_locations_map,
            'organizations': all_organizations
        }, status=status.HTTP_200_OK)



class TraineeAssessmentViewSet(viewsets.ModelViewSet):
    queryset = TraineeAssessment.objects.all().order_by('-created_at')
    serializer_class = TraineeAssessmentSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        queryset = super().get_queryset()
        user = self.request.user
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            queryset = queryset.filter(department__in=user.department_admin.departments)

        code = self.request.query_params.get('code') or self.request.query_params.get('trainee_code')
        if code:
            queryset = queryset.filter(trainee_code__iexact=code.strip())
        dept = self.request.query_params.get('department')
        if dept:
            queryset = queryset.filter(department__iexact=dept.strip())
        status_filter = self.request.query_params.get('status')
        if status_filter:
            queryset = queryset.filter(status__iexact=status_filter.strip())
        search = self.request.query_params.get('search')
        if search:
            s = search.strip()
            queryset = queryset.filter(
                models.Q(trainee_code__icontains=s) |
                models.Q(name__icontains=s) |
                models.Q(department__icontains=s) |
                models.Q(designation__icontains=s)
            )
        return queryset

    @action(detail=False, methods=['get'], url_path='fetch-trainee')
    def fetch_trainee(self, request):
        code = request.query_params.get('code')
        if not code:
            return Response({"error": "Trainee code parameter is required."}, status=status.HTTP_400_BAD_REQUEST)

        code = code.strip().strip('<>').strip()
        user = request.user

        # First check if an assessment already exists
        existing = TraineeAssessment.objects.filter(trainee_code__iexact=code).first()
        if existing:
            # Department admin permission check
            if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
                if existing.department not in user.department_admin.departments:
                    return Response({
                        "error": f"You do not have permission to view assessments in department '{existing.department}'."
                    }, status=status.HTTP_403_FORBIDDEN)

            serializer = self.get_serializer(existing)
            return Response({
                "exists": True,
                "assessment": serializer.data
            }, status=status.HTTP_200_OK)

        # Look up trainee in TrainingEmployee, or fallback to Employee master
        trainee = TrainingEmployee.objects.filter(employee_code__iexact=code).first()
        if not trainee:
            emp = Employee.objects.filter(employee_code__iexact=code).first()
            if emp:
                trainee = TrainingEmployee.objects.create(
                    employee_code=emp.employee_code,
                    name=emp.name,
                    department=emp.department,
                    location=emp.location,
                    branch=emp.branch,
                    designation=emp.designation,
                    date_of_joining=emp.date_of_joining,
                    training_period_months=1,
                    status='Training'
                )
            else:
                return Response({"error": f"Trainee with code '{code}' not found in Training or Employee records."}, status=status.HTTP_404_NOT_FOUND)

        # Department admin authorization check
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            if trainee.department not in user.department_admin.departments:
                return Response({
                    "error": f"You do not have permission to assess trainees in department '{trainee.department}'."
                }, status=status.HTTP_403_FORBIDDEN)

        # Return auto-populated trainee metadata
        return Response({
            "exists": False,
            "trainee_data": {
                "trainee_id": trainee.id,
                "trainee_code": trainee.employee_code,
                "name": trainee.name,
                "department": trainee.department,
                "location": trainee.location or "",
                "branch": trainee.branch.name if trainee.branch else "",
                "designation": trainee.designation,
                "joining_date": trainee.date_of_joining.strftime('%Y-%m-%d') if trainee.date_of_joining else "",
                "training_period": f"{trainee.training_period_months} Month" if trainee.training_period_months == 1 else f"{trainee.training_period_months} Months",
                "training_completion_date": trainee.training_end_date.strftime('%Y-%m-%d') if trainee.training_end_date else "",
                "status": trainee.status
            }
        }, status=status.HTTP_200_OK)

    def create(self, request, *args, **kwargs):
        trainee_id = request.data.get('trainee')
        trainee_code = request.data.get('trainee_code')

        trainee = None
        if trainee_id:
            trainee = TrainingEmployee.objects.filter(id=trainee_id).first()
        elif trainee_code:
            trainee = TrainingEmployee.objects.filter(employee_code__iexact=trainee_code).first()

        if not trainee:
            return Response({"error": "Valid trainee record is required."}, status=status.HTTP_400_BAD_REQUEST)

        # Department admin authorization check
        user = request.user
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            if trainee.department not in user.department_admin.departments:
                return Response({
                    "error": f"You do not have permission to submit assessments for department '{trainee.department}'."
                }, status=status.HTTP_403_FORBIDDEN)

        # Upsert: if assessment exists for this trainee, update it
        existing = TraineeAssessment.objects.filter(trainee=trainee).first()
        data = request.data.copy()

        # Status logic: if explicitly 'Draft', keep 'Draft'; otherwise 'Pending'
        req_status = data.get('status')
        if req_status != 'Draft':
            data['status'] = 'Pending'

        if existing:
            if existing.status == 'Approved' and not user.is_superuser:
                return Response({"error": "Approved assessments are locked and cannot be modified."}, status=status.HTTP_400_BAD_REQUEST)
            serializer = self.get_serializer(existing, data=data, partial=True)
        else:
            data['trainee'] = trainee.id
            if not data.get('trainee_code'):
                data['trainee_code'] = trainee.employee_code
            if not data.get('name'):
                data['name'] = trainee.name
            if not data.get('department'):
                data['department'] = trainee.department
            if not data.get('location'):
                data['location'] = trainee.location
            if not data.get('designation'):
                data['designation'] = trainee.designation
            if not data.get('joining_date') and trainee.date_of_joining:
                data['joining_date'] = trainee.date_of_joining
            if not data.get('training_completion_date') and trainee.training_end_date:
                data['training_completion_date'] = trainee.training_end_date
            serializer = self.get_serializer(data=data)

        if serializer.is_valid():
            assessment = serializer.save(evaluator=request.user)
            return Response(self.get_serializer(assessment).data, status=status.HTTP_201_CREATED if not existing else status.HTTP_200_OK)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], url_path='approve')
    def approve(self, request, pk=None):
        assessment = self.get_object()
        user = request.user

        # Department admin authorization check
        if not user.is_superuser:
            if not hasattr(user, 'department_admin') or assessment.department not in user.department_admin.departments:
                return Response(
                    {"error": f"You do not have permission to approve assessments for department '{assessment.department}'."},
                    status=status.HTTP_403_FORBIDDEN
                )

        from django.utils import timezone
        assessment.status = 'Approved'
        assessment.approved_by = user
        assessment.approved_at = timezone.now()
        remarks = request.data.get('remarks')
        if remarks:
            assessment.approval_remarks = remarks
        assessment.save()

        # Update training employee record
        if assessment.trainee:
            assessment.trainee.status = 'Completed'
            assessment.trainee.save(update_fields=['status'])

        return Response({
            "message": f"Assessment for {assessment.name} ({assessment.trainee_code}) approved successfully!",
            "assessment": self.get_serializer(assessment).data
        }, status=status.HTTP_200_OK)


class IncentivePeriodViewSet(viewsets.ModelViewSet):
    queryset = IncentivePeriod.objects.all()
    serializer_class = IncentivePeriodSerializer
    permission_classes = [IsAdminUserOrReadOnly]

    def get_queryset(self):
        qs = IncentivePeriod.objects.all()
        branch = self.request.query_params.get('branch')
        if branch:
            qs = qs.filter(Q(branch__name__iexact=branch) | Q(branch__id=branch if str(branch).isdigit() else -1))
        return qs

    def perform_create(self, serializer):
        branch_param = self.request.data.get('branch') or self.request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
        if branch_obj and not serializer.validated_data.get('branch'):
            instance = serializer.save(branch=branch_obj)
        else:
            instance = serializer.save()

        if instance.is_active:
            qs = IncentivePeriod.objects.exclude(pk=instance.pk)
            if instance.branch:
                qs = qs.filter(branch=instance.branch)
            else:
                qs = qs.filter(branch__isnull=True)
            qs.update(is_active=False)

    def perform_update(self, serializer):
        branch_param = self.request.data.get('branch') or self.request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
        if branch_obj and not serializer.validated_data.get('branch'):
            instance = serializer.save(branch=branch_obj)
        else:
            instance = serializer.save()

        if instance.is_active:
            qs = IncentivePeriod.objects.exclude(pk=instance.pk)
            if instance.branch:
                qs = qs.filter(branch=instance.branch)
            else:
                qs = qs.filter(branch__isnull=True)
            qs.update(is_active=False)

    @action(detail=True, methods=['post'], url_path='set-active')
    def set_active(self, request, pk=None):
        period = self.get_object()
        qs = IncentivePeriod.objects.exclude(pk=period.pk)
        if period.branch:
            qs = qs.filter(branch=period.branch)
        else:
            qs = qs.filter(branch__isnull=True)
        qs.update(is_active=False)
        period.is_active = True
        period.save(update_fields=['is_active'])
        return Response({"message": f"Period '{period.name}' set as active.", "period": self.get_serializer(period).data})

    def perform_destroy(self, instance):
        # 1. Delete all employees belonging to this period in Employee Master (both active and trashed)
        emp_filter = Q(incentive_period=instance) | Q(assignment_period__iexact=instance.name) | Q(assignment_period__iexact=instance.months)
        if instance.branch:
            emp_filter = emp_filter & Q(branch=instance.branch)
        Employee.objects.filter(emp_filter).delete()

        # 2. Delete all incentive records belonging to this period
        IncentiveRecord.objects.filter(period=instance).delete()

        # 3. Delete all appraisals associated with this period
        app_filter = Q(assignment_period__iexact=instance.name) | Q(assignment_period__iexact=instance.months)
        if instance.branch:
            app_filter = app_filter & Q(branch=instance.branch)
        Appraisal.objects.filter(app_filter).delete()

        # NOTE: DepartmentAdmin records, tagged departments, tagged locations, tagged branches,
        # and admin users are NEVER touched or modified!
        instance.delete()


class DivisionRateViewSet(viewsets.ModelViewSet):
    queryset = DivisionRate.objects.all()
    serializer_class = DivisionRateSerializer
    permission_classes = [IsAdminUserOrReadOnly]

    def get_queryset(self):
        qs = DivisionRate.objects.all()
        branch = self.request.query_params.get('branch')
        if branch:
            qs = qs.filter(Q(branch__name__iexact=branch) | Q(branch__id=branch if str(branch).isdigit() else -1))
        return qs

    def perform_create(self, serializer):
        branch_param = self.request.data.get('branch') or self.request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
        if branch_obj and not serializer.validated_data.get('branch'):
            serializer.save(branch=branch_obj)
        else:
            serializer.save()

    def perform_update(self, serializer):
        branch_param = self.request.data.get('branch') or self.request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
        if branch_obj and not serializer.validated_data.get('branch'):
            serializer.save(branch=branch_obj)
        else:
            serializer.save()

    @action(detail=False, methods=['post'], url_path='sync-from-employees')
    def sync_from_employees(self, request):
        branch_param = request.data.get('branch') or request.query_params.get('branch')
        branch_obj = None
        emp_qs = Employee.objects.exclude(division='')
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
            if branch_obj:
                emp_qs = emp_qs.filter(branch=branch_obj)
        divisions = emp_qs.values_list('division', flat=True).distinct()
        created_count = 0
        for d in divisions:
            d_clean = (d or '').strip()
            if d_clean:
                filter_kwargs = {'division_name': d_clean, 'branch': branch_obj}
                _, created = DivisionRate.objects.get_or_create(defaults={'rate': 0.0}, **filter_kwargs)
                if created:
                    created_count += 1
        return Response({
            "message": f"Synced unique divisions. {created_count} new division(s) added.",
            "divisions": self.get_serializer(self.get_queryset(), many=True).data
        })

    @action(detail=False, methods=['post'], url_path='bulk-update')
    def bulk_update_rates(self, request):
        items = request.data.get('items', [])
        branch_param = request.data.get('branch') or request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
        updated = 0
        for item in items:
            div_id = item.get('id')
            rate_val = item.get('rate')
            div_name = item.get('division_name')
            if div_id is not None and rate_val is not None:
                try:
                    obj = DivisionRate.objects.get(id=div_id)
                    obj.rate = float(rate_val)
                    if branch_obj and not obj.branch:
                        obj.branch = branch_obj
                    obj.save(update_fields=['rate', 'branch', 'updated_at'])
                    updated += 1
                except DivisionRate.DoesNotExist:
                    pass
            elif div_name and rate_val is not None and branch_obj:
                obj, _ = DivisionRate.objects.get_or_create(branch=branch_obj, division_name=div_name, defaults={'rate': float(rate_val)})
                obj.rate = float(rate_val)
                obj.save(update_fields=['rate', 'updated_at'])
                updated += 1
        return Response({
            "message": f"Updated {updated} division rate(s).",
            "divisions": self.get_serializer(self.get_queryset(), many=True).data
        })


class EmployeeLevelConfigViewSet(viewsets.ModelViewSet):
    queryset = EmployeeLevelConfig.objects.all()
    serializer_class = EmployeeLevelConfigSerializer
    permission_classes = [IsAdminUserOrReadOnly]

    def get_queryset(self):
        qs = EmployeeLevelConfig.objects.all()
        branch = self.request.query_params.get('branch')
        if branch:
            qs = qs.filter(Q(branch__name__iexact=branch) | Q(branch__id=branch if str(branch).isdigit() else -1))
        return qs

    def perform_create(self, serializer):
        branch_param = self.request.data.get('branch') or self.request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
        if branch_obj and not serializer.validated_data.get('branch'):
            serializer.save(branch=branch_obj)
        else:
            serializer.save()

    def perform_update(self, serializer):
        branch_param = self.request.data.get('branch') or self.request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
        if branch_obj and not serializer.validated_data.get('branch'):
            serializer.save(branch=branch_obj)
        else:
            serializer.save()

    @action(detail=False, methods=['post'], url_path='sync-from-employees')
    def sync_from_employees(self, request):
        branch_param = request.data.get('branch') or request.query_params.get('branch')
        branch_obj = None
        emp_qs = Employee.objects.exclude(level='')
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
            if branch_obj:
                emp_qs = emp_qs.filter(branch=branch_obj)
        levels = emp_qs.values_list('level', flat=True).distinct()
        created_count = 0
        for lvl in levels:
            lvl_clean = (lvl or '').strip().upper()
            if lvl_clean:
                filter_kwargs = {'level_code': lvl_clean, 'branch': branch_obj}
                _, created = EmployeeLevelConfig.objects.get_or_create(
                    defaults={'description': f'Level {lvl_clean}', 'score': 1.0},
                    **filter_kwargs
                )
                if created:
                    created_count += 1
        return Response({
            "message": f"Synced unique employee levels. {created_count} new level(s) added.",
            "levels": self.get_serializer(self.get_queryset(), many=True).data
        })


class ExperienceLevelConfigViewSet(viewsets.ModelViewSet):
    queryset = ExperienceLevelConfig.objects.all()
    serializer_class = ExperienceLevelConfigSerializer
    permission_classes = [IsAdminUserOrReadOnly]

    def get_queryset(self):
        qs = ExperienceLevelConfig.objects.all()
        branch = self.request.query_params.get('branch')
        if branch:
            qs = qs.filter(Q(branch__name__iexact=branch) | Q(branch__id=branch if str(branch).isdigit() else -1))
        return qs

    def perform_create(self, serializer):
        branch_param = self.request.data.get('branch') or self.request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
        if branch_obj and not serializer.validated_data.get('branch'):
            serializer.save(branch=branch_obj)
        else:
            serializer.save()

    def perform_update(self, serializer):
        branch_param = self.request.data.get('branch') or self.request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()
        if branch_obj and not serializer.validated_data.get('branch'):
            serializer.save(branch=branch_obj)
        else:
            serializer.save()


def compute_employee_incentive(emp, period, div_rates_map, level_scores_map, exp_brackets, appraisal=None, existing_record=None):
    # Cut-off date & Experience
    cutoff_date = period.cutoff_date if period else datetime.date.today()
    doj = emp.date_of_joining or datetime.date.today()

    exp_months = 0
    if cutoff_date and doj:
        exp_months = (cutoff_date.year - doj.year) * 12 + (cutoff_date.month - doj.month)
        if cutoff_date.day < doj.day:
            exp_months = max(0, exp_months - 1)
        if exp_months < 0:
            exp_months = 0

    exp_years = exp_months // 12
    rem_months = exp_months % 12
    exp_display = f"{exp_years} yrs {rem_months} mos ({exp_months}m)" if exp_years > 0 else f"{exp_months} mos"

    # Level 1 Score (Employee Level, e.g. MAN, SUP, LOW, MID)
    emp_lvl_clean = (emp.level or '').strip().upper()
    level_1_score = level_scores_map.get(emp_lvl_clean, level_scores_map.get(emp_lvl_clean.lower(), 1.0))
    if not emp_lvl_clean:
        level_1_score = 1.0

    # Level 2 Score (Experience brackets)
    level_2_score = 1.0
    matched_exp_bracket = None
    for bracket in exp_brackets:
        if bracket.min_months <= exp_months <= bracket.max_months:
            level_2_score = bracket.score
            matched_exp_bracket = bracket.name
            break

    # Level 3 Score (Appraisal Assessment Score)
    # Score basis: (marks * 3) / 100
    # e.g. 100 marks -> 3.0, 99 marks -> 2.97
    appraisal_marks = 0.0
    level_3_score = 0.0
    appraisal_id = None
    if appraisal:
        appraisal_id = appraisal.id
        appraisal_marks = float(appraisal.marks_obtained if appraisal.marks_obtained is not None and appraisal.marks_obtained > 0 else max(0, appraisal.performance_score - appraisal.total_deduction))
        level_3_score = round((appraisal_marks * 3.0) / 100.0, 2)

    # Division Rate
    emp_div_clean = (emp.division or '').strip().lower()
    division_rate = div_rates_map.get(emp_div_clean, 0.0)

    # Total Score & Final Amount: (Level 1 + Level 2 + Level 3) * Division Rate
    total_unit_score = round(float(level_1_score) + float(level_2_score) + float(level_3_score), 2)
    final_incentive_amount = round(total_unit_score * float(division_rate), 2)

    is_processed = existing_record and existing_record.status in ['Processed', 'Approved']
    rec_status = existing_record.status if existing_record else 'Draft'

    return {
        'id': existing_record.id if existing_record else f"temp_{emp.id}",
        'record_id': existing_record.id if existing_record else None,
        'employee_id': emp.id,
        'employee_code': emp.employee_code,
        'employee_name': emp.name,
        'department': emp.department,
        'location': emp.location or '',
        'branch': emp.branch.name if emp.branch else '',
        'designation': emp.designation,
        'division': emp.division or '',
        'category': emp.category or '',
        'organization': emp.organization or '',
        # Level 1
        'employee_level': emp.level or 'Standard',
        'level_code': emp.level or 'Standard',
        'level_1_score': level_1_score,
        'level1_score': level_1_score,
        # Level 2
        'joining_date': doj.strftime('%Y-%m-%d') if doj else '',
        'date_of_joining': doj.strftime('%Y-%m-%d') if doj else '',
        'cutoff_date': cutoff_date.strftime('%Y-%m-%d') if cutoff_date else '',
        'experience_months': exp_months,
        'months_of_experience': exp_months,
        'experience_display': exp_display,
        'matched_exp_bracket': matched_exp_bracket,
        'experience_bracket': matched_exp_bracket,
        'level_2_score': level_2_score,
        'level2_score': level_2_score,
        # Level 3
        'appraisal_id': appraisal_id,
        'appraisal_marks': appraisal_marks,
        'marks_obtained': appraisal_marks,
        'level_3_score': level_3_score,
        'level3_score': level_3_score,
        'appraisal_status': appraisal.status if appraisal else 'No Appraisal',
        # Summary & Payout
        'total_unit_score': total_unit_score,
        'total_units': total_unit_score,
        'division_rate': division_rate,
        'final_incentive_amount': final_incentive_amount,
        'final_amount': final_incentive_amount,
        'status': rec_status,
        'is_processed': is_processed,
        'processed_at': existing_record.processed_at.strftime('%Y-%m-%d %H:%M') if existing_record and existing_record.processed_at else '',
        'processed_by': existing_record.processed_by.username if existing_record and existing_record.processed_by else ''
    }


class IncentiveCalculationView(APIView):
    permission_classes = [IsAdminUserOrReadOnly]

    def get(self, request):
        user = request.user
        period_id = request.query_params.get('period_id')
        division_param = request.query_params.get('division')
        dept_param = request.query_params.get('department')
        loc_param = request.query_params.get('location')
        branch_param = request.query_params.get('branch')
        status_param = request.query_params.get('status')
        search_query = request.query_params.get('search')

        # Select branch
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()

        # Select period
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            active_p_qs = IncentivePeriod.objects.filter(is_active=True)
            if branch_obj:
                active_p_qs = active_p_qs.filter(branch=branch_obj)
            period = active_p_qs.first() or IncentivePeriod.objects.filter(is_active=True).first() or IncentivePeriod.objects.first()
        elif period_id:
            period = IncentivePeriod.objects.filter(id=period_id).first()
        else:
            active_p_qs = IncentivePeriod.objects.filter(is_active=True)
            if branch_obj:
                active_p_qs = active_p_qs.filter(branch=branch_obj)
            period = active_p_qs.first() or IncentivePeriod.objects.filter(is_active=True).first() or IncentivePeriod.objects.first()

        if not period:
            return Response({"error": "No incentive period defined. Please create an incentive period first."}, status=status.HTTP_404_NOT_FOUND)

        # Division rates & Level configurations scoped to the branch
        div_rates_qs = DivisionRate.objects.all()
        lvl_qs = EmployeeLevelConfig.objects.all()
        exp_qs = ExperienceLevelConfig.objects.all()
        if branch_obj:
            div_rates_qs = div_rates_qs.filter(branch=branch_obj)
            lvl_qs = lvl_qs.filter(branch=branch_obj)
            exp_qs = exp_qs.filter(branch=branch_obj)

        div_rates_map = {d.division_name.strip().lower(): d.rate for d in div_rates_qs}
        level_scores_map = {l.level_code.strip().upper(): l.score for l in lvl_qs}
        for l in lvl_qs:
            level_scores_map[l.level_code.strip().lower()] = l.score
        exp_brackets = list(exp_qs.order_by('min_months'))

        # Employees query - strictly isolated to the active or selected period, excluding trashed
        emp_qs = Employee.objects.filter(status='Active', is_trashed=False).filter(
            Q(incentive_period=period) | 
            Q(incentive_records__period=period) | 
            Q(assignment_period__iexact=period.name) | 
            Q(assignment_period__iexact=period.months)
        ).distinct()

        # Department admin restrictions
        if user and user.is_authenticated and not user.is_superuser and hasattr(user, 'department_admin'):
            dept_admin = user.department_admin
            if dept_admin.departments:
                dept_filter = Q()
                for d in dept_admin.departments:
                    dept_filter |= Q(department__iexact=d) | Q(department__icontains=d)
                emp_qs = emp_qs.filter(dept_filter)
            if dept_admin.locations:
                loc_filter = Q()
                for loc in dept_admin.locations:
                    loc_filter |= build_exact_location_filter(loc)
                emp_qs = emp_qs.filter(loc_filter)
            if dept_admin.branches:
                br_filter = Q()
                for br in dept_admin.branches:
                    br_filter |= Q(branch__name__iexact=br)
                emp_qs = emp_qs.filter(br_filter)

        # Filters
        if division_param:
            emp_qs = emp_qs.filter(division__iexact=division_param)
        if dept_param:
            emp_qs = emp_qs.filter(department__iexact=dept_param)
        if loc_param:
            emp_qs = emp_qs.filter(build_exact_location_filter(loc_param))
        if branch_param:
            emp_qs = emp_qs.filter(branch__name__iexact=branch_param)
        if search_query:
            emp_qs = emp_qs.filter(Q(employee_code__icontains=search_query) | Q(name__icontains=search_query))

        # Fetch appraisals matching this period or fallback to latest
        appraisals_map = {}
        for app in Appraisal.objects.filter(Q(assignment_period__iexact=period.name) | Q(assignment_period__iexact=period.months)).order_by('-submitted_date'):
            if app.employee_code not in appraisals_map:
                appraisals_map[app.employee_code] = app
        for app in Appraisal.objects.order_by('-submitted_date'):
            if app.employee_code not in appraisals_map:
                appraisals_map[app.employee_code] = app

        # Fetch existing records for this period
        records_map = {r.employee_id: r for r in IncentiveRecord.objects.filter(period=period)}

        results = []
        for emp in emp_qs.order_by('employee_code'):
            appraisal = appraisals_map.get(emp.employee_code)
            record = records_map.get(emp.id)
            item = compute_employee_incentive(emp, period, div_rates_map, level_scores_map, exp_brackets, appraisal, record)
            
            # Status filter
            if status_param:
                if status_param.lower() == 'processed' and not item['is_processed']:
                    continue
                elif status_param.lower() == 'draft' and item['is_processed']:
                    continue

            results.append(item)

        total_amount = sum(r['final_incentive_amount'] for r in results)
        processed_count = sum(1 for r in results if r['is_processed'])

        return Response({
            'period': IncentivePeriodSerializer(period).data,
            'summary': {
                'total_employees': len(results),
                'total_incentive_amount': round(total_amount, 2),
                'processed_count': processed_count,
                'draft_count': len(results) - processed_count
            },
            'records': results
        })


class IncentiveProcessActionView(APIView):
    permission_classes = [IsAdminUserOrReadOnly]

    def post(self, request):
        user = request.user
        period_id = request.data.get('period_id')
        employee_ids = request.data.get('employee_ids', [])
        select_all = request.data.get('select_all', False)
        action_type = request.data.get('action', 'process') # 'process' or 'revert'

        if not period_id:
            return Response({"error": "Period ID is required."}, status=status.HTTP_400_BAD_REQUEST)

        period = IncentivePeriod.objects.filter(id=period_id).first()
        if not period:
            return Response({"error": "Incentive period not found."}, status=status.HTTP_404_NOT_FOUND)

        if select_all:
            emp_qs = Employee.objects.filter(status='Active').filter(
                Q(incentive_period=period) | 
                Q(incentive_records__period=period) | 
                Q(assignment_period__iexact=period.name) | 
                Q(assignment_period__iexact=period.months)
            ).distinct()
            # Restrict if department admin
            if user and not user.is_superuser and hasattr(user, 'department_admin'):
                dept_admin = user.department_admin
                if dept_admin.departments:
                    dept_filter = Q()
                    for d in dept_admin.departments:
                        dept_filter |= Q(department__iexact=d) | Q(department__icontains=d)
                    emp_qs = emp_qs.filter(dept_filter)
        else:
            if not employee_ids:
                return Response({"error": "Please provide employee IDs to process."}, status=status.HTTP_400_BAD_REQUEST)
            emp_qs = Employee.objects.filter(id__in=employee_ids)

        from django.utils import timezone

        # Branch-scoped configurations cache
        branch_configs_cache = {}
        def get_branch_configs_for_emp(b_obj):
            bid = b_obj.id if b_obj else 'null'
            if bid in branch_configs_cache:
                return branch_configs_cache[bid]
            
            div_qs = DivisionRate.objects.filter(branch=b_obj) if b_obj else DivisionRate.objects.filter(branch__isnull=True)
            lvl_qs = EmployeeLevelConfig.objects.filter(branch=b_obj) if b_obj else EmployeeLevelConfig.objects.filter(branch__isnull=True)
            exp_qs = ExperienceLevelConfig.objects.filter(branch=b_obj) if b_obj else ExperienceLevelConfig.objects.filter(branch__isnull=True)

            # Fallback for levels/experience if branch hasn't configured custom brackets yet
            if not lvl_qs.exists():
                lvl_qs = EmployeeLevelConfig.objects.all()
            if not exp_qs.exists():
                exp_qs = ExperienceLevelConfig.objects.all()

            d_map = {d.division_name.strip().lower(): d.rate for d in div_qs}
            l_map = {l.level_code.strip().upper(): l.score for l in lvl_qs}
            for l in lvl_qs:
                l_map[l.level_code.strip().lower()] = l.score
            e_brackets = list(exp_qs.order_by('min_months'))
            branch_configs_cache[bid] = (d_map, l_map, e_brackets)
            return branch_configs_cache[bid]

        appraisals_map = {}
        for app in Appraisal.objects.order_by('-submitted_date'):
            if app.employee_code not in appraisals_map:
                appraisals_map[app.employee_code] = app

        count = 0
        for emp in emp_qs:
            appraisal = appraisals_map.get(emp.employee_code)
            rec, _ = IncentiveRecord.objects.get_or_create(period=period, employee=emp)
            
            emp_div_map, emp_lvl_map, emp_exp_brackets = get_branch_configs_for_emp(emp.branch)
            computed = compute_employee_incentive(emp, period, emp_div_map, emp_lvl_map, emp_exp_brackets, appraisal, rec)

            rec.employee_code = emp.employee_code
            rec.employee_name = emp.name
            rec.department = emp.department
            rec.location = emp.location or ''
            rec.branch = emp.branch
            rec.designation = emp.designation
            rec.division = emp.division or ''
            rec.category = emp.category or ''
            rec.employee_level = computed['employee_level']
            rec.level_1_score = computed['level_1_score']
            rec.joining_date = emp.date_of_joining
            rec.cutoff_date = period.cutoff_date
            rec.experience_months = computed['experience_months']
            rec.experience_display = computed['experience_display']
            rec.level_2_score = computed['level_2_score']
            rec.appraisal = appraisal
            rec.appraisal_marks = computed['appraisal_marks']
            rec.level_3_score = computed['level_3_score']
            rec.division_rate = computed['division_rate']

            if action_type == 'process':
                rec.status = 'Processed'
                rec.processed_by = user
                rec.processed_at = timezone.now()
            else:
                rec.status = 'Draft'
                rec.processed_by = None
                rec.processed_at = None

            rec.save()
            count += 1

        action_word = "processed" if action_type == 'process' else "reverted"
        return Response({
            "message": f"Successfully {action_word} incentives for {count} staff record(s).",
            "count": count
        }, status=status.HTTP_200_OK)


class IncentiveExcelExportView(APIView):
    permission_classes = [IsAdminUserOrReadOnly]

    def get(self, request):
        user = request.user
        period_id = request.query_params.get('period_id')
        division_param = request.query_params.get('division')
        dept_param = request.query_params.get('department')
        status_param = request.query_params.get('status')

        branch_param = request.query_params.get('branch')
        branch_obj = None
        if branch_param:
            branch_obj = Branch.objects.filter(Q(name__iexact=branch_param) | Q(id=branch_param if str(branch_param).isdigit() else -1)).first()

        if period_id:
            period = IncentivePeriod.objects.filter(id=period_id).first()
        else:
            active_p_qs = IncentivePeriod.objects.filter(is_active=True)
            if branch_obj:
                active_p_qs = active_p_qs.filter(branch=branch_obj)
            period = active_p_qs.first() or IncentivePeriod.objects.filter(is_active=True).first() or IncentivePeriod.objects.first()

        if not period:
            return Response({"error": "No incentive period found."}, status=status.HTTP_404_NOT_FOUND)

        div_rates_qs = DivisionRate.objects.all()
        lvl_qs = EmployeeLevelConfig.objects.all()
        exp_qs = ExperienceLevelConfig.objects.all()
        if branch_obj:
            div_rates_qs = div_rates_qs.filter(branch=branch_obj)
            lvl_qs = lvl_qs.filter(branch=branch_obj)
            exp_qs = exp_qs.filter(branch=branch_obj)

        div_rates_map = {d.division_name.strip().lower(): d.rate for d in div_rates_qs}
        level_scores_map = {l.level_code.strip().upper(): l.score for l in lvl_qs}
        for l in lvl_qs:
            level_scores_map[l.level_code.strip().lower()] = l.score
        exp_brackets = list(exp_qs.order_by('min_months'))

        emp_qs = Employee.objects.filter(status='Active', is_trashed=False).filter(
            Q(incentive_period=period) | 
            Q(incentive_records__period=period) | 
            Q(assignment_period__iexact=period.name) | 
            Q(assignment_period__iexact=period.months)
        ).distinct()
        if branch_obj:
            emp_qs = emp_qs.filter(branch=branch_obj)
        if user and not user.is_superuser and hasattr(user, 'department_admin'):
            dept_admin = user.department_admin
            if dept_admin.departments:
                dept_filter = Q()
                for d in dept_admin.departments:
                    dept_filter |= Q(department__iexact=d) | Q(department__icontains=d)
                emp_qs = emp_qs.filter(dept_filter)

        if division_param:
            emp_qs = emp_qs.filter(division__iexact=division_param)
        if dept_param:
            emp_qs = emp_qs.filter(department__iexact=dept_param)

        appraisals_map = {}
        for app in Appraisal.objects.filter(Q(assignment_period__iexact=period.name) | Q(assignment_period__iexact=period.months)).order_by('-submitted_date'):
            if app.employee_code not in appraisals_map:
                appraisals_map[app.employee_code] = app
        for app in Appraisal.objects.order_by('-submitted_date'):
            if app.employee_code not in appraisals_map:
                appraisals_map[app.employee_code] = app

        records_map = {r.employee_id: r for r in IncentiveRecord.objects.filter(period=period)}

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Incentive Report"

        headers = [
            "Employee Code", "Employee Name", "Division", "Department", "Duty Location",
            "Joining Date", "Cut-off Date", "Experience", 
            "Staff Level (Level 1)", "Level 1 Score", "Level 2 Score (Exp)",
            "Assessment Marks (out of 100)", "Level 3 Score (Assessment)",
            "Total Score Units", "Division Rate (INR)", "Final Incentive Amount (INR)",
            "Status", "Processed Date", "Processed By"
        ]

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

        row_num = 2
        total_payout = 0.0

        for emp in emp_qs.order_by('employee_code'):
            appraisal = appraisals_map.get(emp.employee_code)
            record = records_map.get(emp.id)
            c = compute_employee_incentive(emp, period, div_rates_map, level_scores_map, exp_brackets, appraisal, record)

            if status_param:
                if status_param.lower() == 'processed' and not c['is_processed']:
                    continue
                elif status_param.lower() == 'draft' and c['is_processed']:
                    continue

            ws.cell(row=row_num, column=1, value=c['employee_code'])
            ws.cell(row=row_num, column=2, value=c['employee_name'])
            ws.cell(row=row_num, column=3, value=c['division'])
            ws.cell(row=row_num, column=4, value=c['department'])
            ws.cell(row=row_num, column=5, value=c['location'])
            ws.cell(row=row_num, column=6, value=c['joining_date'])
            ws.cell(row=row_num, column=7, value=c['cutoff_date'])
            ws.cell(row=row_num, column=8, value=c['experience_display'])
            ws.cell(row=row_num, column=9, value=c['employee_level'])
            ws.cell(row=row_num, column=10, value=c['level_1_score'])
            ws.cell(row=row_num, column=11, value=c['level_2_score'])
            ws.cell(row=row_num, column=12, value=c['appraisal_marks'])
            ws.cell(row=row_num, column=13, value=c['level_3_score'])
            ws.cell(row=row_num, column=14, value=c['total_unit_score'])
            ws.cell(row=row_num, column=15, value=c['division_rate'])
            ws.cell(row=row_num, column=16, value=c['final_incentive_amount'])
            ws.cell(row=row_num, column=17, value=c['status'])
            ws.cell(row=row_num, column=18, value=c['processed_at'])
            ws.cell(row=row_num, column=19, value=c['processed_by'])

            total_payout += c['final_incentive_amount']
            row_num += 1

        # Summary Row
        ws.cell(row=row_num, column=15, value="TOTAL PAYOUT:").font = openpyxl.styles.Font(bold=True)
        tot_cell = ws.cell(row=row_num, column=16, value=total_payout)
        tot_cell.font = openpyxl.styles.Font(bold=True, color="166534")

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 12)

        period_title = period.name.replace(' ', '_').replace(',', '')
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f"attachment; filename=incentive_report_{period_title}.xlsx"
        wb.save(response)
        return response


class DutyLocationViewSet(viewsets.ModelViewSet):
    """
    CRUD API for managing Hospital Duty Locations (wards, units, clinics, etc.).
    Supports searching, filtering by department, branch, and status,
    as well as syncing any untracked locations from Employee records.
    """
    queryset = DutyLocation.objects.all().select_related('branch').order_by('name')
    serializer_class = DutyLocationSerializer
    permission_classes = [permissions.IsAuthenticated]

    def check_permissions(self, request):
        super().check_permissions(request)
        user = request.user
        if user and user.is_authenticated and hasattr(user, 'department_admin') and not user.is_superuser:
            self.permission_denied(
                request,
                message="Department Administrators do not have access to Duty Location Master."
            )

    def get_queryset(self):
        user = self.request.user
        if user and user.is_authenticated and hasattr(user, 'department_admin') and not user.is_superuser:
            return DutyLocation.objects.none()

        queryset = super().get_queryset()

        search = self.request.query_params.get('search', '').strip()
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) |
                Q(code__icontains=search) |
                Q(department__icontains=search) |
                Q(description__icontains=search)
            )

        department = self.request.query_params.get('department', '').strip()
        if department:
            queryset = queryset.filter(department__iexact=department)

        branch_id = self.request.query_params.get('branch', '').strip()
        if branch_id:
            if branch_id.isdigit():
                queryset = queryset.filter(branch_id=branch_id)
            else:
                queryset = queryset.filter(branch__name__iexact=branch_id)

        is_active = self.request.query_params.get('is_active', '').strip()
        if is_active.lower() in ['true', '1']:
            queryset = queryset.filter(is_active=True)
        elif is_active.lower() in ['false', '0']:
            queryset = queryset.filter(is_active=False)

        return queryset

    @action(detail=False, methods=['post'])
    def sync_from_employees(self, request):
        """
        Scans Employee and TrainingEmployee tables to find any locations not yet registered in DutyLocation master,
        and auto-creates them.
        """
        emp_locs = Employee.objects.exclude(location='').values_list('location', 'department', 'branch_id').distinct()
        trn_locs = TrainingEmployee.objects.exclude(location='').values_list('location', 'department').distinct()

        existing_names = set(DutyLocation.objects.values_list('name', flat=True))
        existing_names_lower = {n.strip().lower() for n in existing_names if n}

        created_count = 0
        created_names = []

        # Process Employee locations
        for loc_tuple in emp_locs:
            loc_name = loc_tuple[0].strip() if loc_tuple[0] else ''
            dept_name = loc_tuple[1].strip() if loc_tuple[1] else ''
            branch_id = loc_tuple[2]
            if loc_name and loc_name.lower() not in existing_names_lower:
                DutyLocation.objects.create(
                    name=loc_name,
                    department=dept_name,
                    branch_id=branch_id,
                    is_active=True
                )
                existing_names_lower.add(loc_name.lower())
                created_names.append(loc_name)
                created_count += 1

        # Process TrainingEmployee locations
        for trn_tuple in trn_locs:
            loc_name = trn_tuple[0].strip() if trn_tuple[0] else ''
            dept_name = trn_tuple[1].strip() if trn_tuple[1] else ''
            if loc_name and loc_name.lower() not in existing_names_lower:
                DutyLocation.objects.create(
                    name=loc_name,
                    department=dept_name,
                    is_active=True
                )
                existing_names_lower.add(loc_name.lower())
                created_names.append(loc_name)
                created_count += 1

        total_count = DutyLocation.objects.count()
        return Response({
            'message': f"Synced locations. {created_count} new location(s) added.",
            'created_count': created_count,
            'created_locations': created_names,
            'total_locations': total_count
        }, status=status.HTTP_200_OK)

    @action(detail=True, methods=['post'])
    def toggle_status(self, request, pk=None):
        location = self.get_object()
        location.is_active = not location.is_active
        location.save()
        return Response({
            'id': location.id,
            'name': location.name,
            'is_active': location.is_active,
            'message': f"Location '{location.name}' is now {'active' if location.is_active else 'inactive'}."
        }, status=status.HTTP_200_OK)




